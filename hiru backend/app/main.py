
from __future__ import annotations

import json
import logging
import re
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from starlette.concurrency import run_in_threadpool

from app.services.concept_extractor import extract_concepts
from app.services.knowledge_graph import build_knowledge_graph
from app.services.semantic_matcher import match_concepts_to_text
from app.services.transcription import transcribe


# =========================================================
# APPLICATION CONFIGURATION
# =========================================================

app = FastAPI(
    title="EduTwin-AI C1 Backend",
    description=(
        "Lecture document processing, automatic transcription, "
        "knowledge graph generation, and timestamped lecture analysis."
    ),
    version="0.5.1",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://127.0.0.1:5500",
        "http://localhost:5500",
    ],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

# main.py is in backend/app/main.py.
# parents[1] resolves to the backend directory.
BACKEND_DIR = Path(__file__).resolve().parents[1]

DATA_DIR = BACKEND_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
RECORDING_DIR = DATA_DIR / "recordings"
PROCESSED_DIR = DATA_DIR / "processed"
ANNOTATION_DIR = DATA_DIR / "annotations"
MODEL_DIR = BACKEND_DIR / "models"
REPORT_DIR = BACKEND_DIR / "reports"

for directory in (
    UPLOAD_DIR,
    RECORDING_DIR,
    PROCESSED_DIR,
    ANNOTATION_DIR,
    MODEL_DIR,
    REPORT_DIR,
):
    directory.mkdir(parents=True, exist_ok=True)

MAX_DOCUMENT_SIZE = 100 * 1024 * 1024
MAX_RECORDING_SIZE = 200 * 1024 * 1024

SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".pptx"}
SUPPORTED_RECORDING_EXTENSIONS = {
    ".mp4",
    ".mp3",
    ".wav",
    ".m4a",
    ".webm",
    ".mpeg",
    ".mpga",
    ".ogg",
    ".flac",
}

logger = logging.getLogger(__name__)


# =========================================================
# HEALTH ENDPOINTS
# =========================================================

@app.get("/")
def home():
    return {
        "project": "EduTwin-AI",
        "component": "C1",
        "status": "running",
        "version": app.version,
    }


@app.get("/health")
def health():
    return {"status": "healthy"}


# =========================================================
# UPLOAD VALIDATION
# =========================================================

async def read_upload(
    file: UploadFile,
    max_size: int,
) -> bytes:
    """Read an upload while enforcing its configured size limit."""
    content = await file.read(max_size + 1)

    if len(content) > max_size:
        limit_mb = max_size // (1024 * 1024)
        raise HTTPException(
            status_code=413,
            detail=f"File exceeds the {limit_mb} MB limit.",
        )

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    return content


def validate_document(
    filename: str | None,
    content: bytes,
) -> str:
    """Validate supported lecture document types."""
    if not filename:
        raise HTTPException(
            status_code=400,
            detail="The document must have a filename.",
        )

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Please upload a PDF or PowerPoint (.pptx) file.",
        )

    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is not a valid PDF.",
        )

    if extension == ".pptx" and not content.startswith(b"PK"):
        raise HTTPException(
            status_code=400,
            detail="The uploaded file does not appear to be a valid PPTX.",
        )

    return extension


def validate_recording(
    filename: str | None,
) -> str:
    """Validate supported audio and video extensions."""
    if not filename:
        raise HTTPException(
            status_code=400,
            detail="The recording must have a filename.",
        )

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_RECORDING_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported recording format. Supported formats include "
                "MP4, MP3, WAV, M4A, WEBM, MPEG, OGG, and FLAC."
            ),
        )

    return extension


# =========================================================
# PDF AND POWERPOINT TEXT EXTRACTION
# =========================================================

def extract_pdf_pages(reader: PdfReader) -> list[dict]:
    """Extract text from each PDF page."""
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        pages.append({
            "page": page_number,
            "text": text.strip(),
        })

    return pages


def extract_pptx_pages(content: bytes) -> list[dict]:
    """Extract text from each PowerPoint slide."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                "PowerPoint support is missing. Activate the backend "
                "environment and run: python -m pip install python-pptx"
            ),
        ) from exc

    try:
        presentation = Presentation(BytesIO(content))
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not read the PowerPoint presentation.",
        ) from exc

    pages = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text = []

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]

                    if cells:
                        slide_text.append(" | ".join(cells))

        pages.append({
            "page": slide_number,
            "text": "\n".join(slide_text).strip(),
        })

    return pages


def extract_document_pages(
    filename: str,
    content: bytes,
) -> list[dict]:
    """Select the appropriate extractor for PDF or PPTX."""
    extension = validate_document(filename, content)

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDFs are not supported.",
                )

            return extract_pdf_pages(reader)

        return extract_pptx_pages(content)

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Document text extraction failed")

        raise HTTPException(
            status_code=422,
            detail="Could not extract text from the lecture document.",
        ) from exc


def combine_page_text(pages: list[dict]) -> str:
    """Combine non-empty page or slide text."""
    return "\n\n".join(
        f"[Page {page['page']}]\n{page['text']}"
        for page in pages
        if page.get("text")
    ).strip()


# =========================================================
# PDF UPLOAD ENDPOINT
# =========================================================

@app.post("/api/lectures/upload")
async def upload_lecture(file: UploadFile = File(...)):
    """Upload a PDF and return extracted concepts and page text."""
    try:
        content = await read_upload(file, MAX_DOCUMENT_SIZE)
        extension = validate_document(file.filename, content)

        if extension != ".pdf":
            raise HTTPException(
                status_code=400,
                detail="This endpoint accepts PDF files only.",
            )

        pages = extract_document_pages(
            file.filename or "",
            content,
        )
        full_text = combine_page_text(pages)

        if not full_text:
            raise HTTPException(
                status_code=422,
                detail=(
                    "No extractable text was found. "
                    "The PDF may require OCR."
                ),
            )

        lecture_id = str(uuid4())
        saved_path = UPLOAD_DIR / f"{lecture_id}.pdf"
        saved_path.write_bytes(content)

        concepts = extract_concepts(full_text)

        return {
            "lecture_id": lecture_id,
            "filename": file.filename,
            "page_count": len(pages),
            "extracted_character_count": len(full_text),
            "concept_count": len(concepts),
            "concepts": concepts,
            "pages": pages,
            "message": "PDF processed successfully.",
        }

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("PDF upload failed")

        raise HTTPException(
            status_code=500,
            detail="Could not process the uploaded PDF.",
        ) from exc
    finally:
        await file.close()


# =========================================================
# KNOWLEDGE GRAPH ENDPOINT
# =========================================================

@app.post("/api/lectures/graph")
async def generate_lecture_graph(file: UploadFile = File(...)):
    """Generate a candidate knowledge graph from a PDF or PPTX."""
    try:
        content = await read_upload(file, MAX_DOCUMENT_SIZE)

        pages = extract_document_pages(
            file.filename or "",
            content,
        )
        full_text = combine_page_text(pages)

        if not full_text:
            raise HTTPException(
                status_code=422,
                detail="No extractable text was found in the document.",
            )

        concepts = extract_concepts(full_text)
        graph = build_knowledge_graph(pages, concepts)

        return {
            "lecture_id": str(uuid4()),
            "filename": file.filename,
            "page_count": len(pages),
            "concept_count": len(concepts),
            "concepts": concepts,
            "graph": graph,
            "message": "Candidate knowledge graph generated successfully.",
        }

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Knowledge graph generation failed")

        raise HTTPException(
            status_code=500,
            detail="Could not generate the knowledge graph.",
        ) from exc
    finally:
        await file.close()


# =========================================================
# TRANSCRIPTION NORMALIZATION
# =========================================================

def normalize_transcription_result(result) -> dict:
    """Normalize common faster-whisper transcription results."""
    if isinstance(result, dict):
        data = result
    elif hasattr(result, "model_dump"):
        data = result.model_dump()
    elif hasattr(result, "__dict__"):
        data = vars(result)
    else:
        raise ValueError(
            "The transcription service returned an unsupported result."
        )

    raw_segments = data.get("segments", [])
    normalized_segments = []

    for index, segment in enumerate(raw_segments):
        if hasattr(segment, "model_dump"):
            segment = segment.model_dump()
        elif not isinstance(segment, dict):
            segment = vars(segment)

        normalized_segments.append({
            "id": segment.get("id", index),
            "start": float(segment.get("start", 0) or 0),
            "end": float(segment.get("end", 0) or 0),
            "text": str(segment.get("text", "")).strip(),
        })

    text = str(data.get("text", "")).strip()

    if not text:
        text = " ".join(
            segment["text"]
            for segment in normalized_segments
            if segment["text"]
        )

    return {
        "segments": normalized_segments,
        "language": data.get("language"),
        "text": text,
    }


# =========================================================
# TIMESTAMPED GRAPH TIMELINE
# =========================================================

def build_timeline_compatibility(
    transcript_data: dict,
    pages: list[dict],
    similarity_threshold: float = 0.45,
) -> dict:
    """
    Build cumulative graph snapshots by matching timestamped transcript
    segments against curriculum concepts.

    MiniLM semantic matching is attempted first. If the pretrained model
    cannot be loaded, the matcher uses exact phrase matching. A concept match
    is evidence of a possible mention, not proof of mastery or full coverage.
    """
    concepts = extract_concepts(combine_page_text(pages))
    initial_graph = build_knowledge_graph(pages, concepts)

    segments = transcript_data.get("segments", [])
    timeline = []
    evidence_counts = {
        str(node["id"]): 0
        for node in initial_graph.get("nodes", [])
    }

    for index, segment in enumerate(segments):
        text = str(segment.get("text", "")).strip()
        start = float(segment.get("start", 0) or 0)
        end = float(segment.get("end", start) or start)

        match_result = match_concepts_to_text(
            text,
            concepts,
            threshold=similarity_threshold,
        )
        match_details = match_result.get("matches", [])
        matched_concepts = [
            str(match["name"]).casefold()
            for match in match_details
        ]
        matched_scores = {
            str(match["name"]).casefold(): float(match.get("similarity", 0))
            for match in match_details
        }

        snapshot_nodes = []
        for node in initial_graph.get("nodes", []):
            node_id = str(node["id"]).casefold()
            if node_id in matched_concepts:
                evidence_counts[str(node["id"])] = (
                    evidence_counts.get(str(node["id"]), 0) + 1
                )

            count = evidence_counts.get(str(node["id"]), 0)
            snapshot_nodes.append({
                **node,
                "evidence_count": count,
                "status": "introduced" if count > 0 else "not_yet_observed",
                "latest_similarity": matched_scores.get(node_id),
            })

        timeline.append({
            "segment_id": segment.get("id", index),
            "start_seconds": start,
            "end_seconds": end,
            "transcript_text": text,
            "matched_concepts": [
                match["name"] for match in match_details
            ],
            "concept_matches": match_details,
            "matching_method": match_result.get("method", "none"),
            "nodes": snapshot_nodes,
            "edges": initial_graph.get("edges", []),
        })

    return {
        "initial_graph": {
            **initial_graph,
            "nodes": [
                {
                    **node,
                    "evidence_count": 0,
                    "status": "not_yet_observed",
                    "latest_similarity": None,
                }
                for node in initial_graph.get("nodes", [])
            ],
        },
        "timeline": timeline,
        "matching": {
            "preferred_method": "MiniLM cosine similarity",
            "model": "sentence-transformers/all-MiniLM-L6-v2",
            "similarity_threshold": similarity_threshold,
            "fallback": "case-insensitive exact phrase matching",
        },
        "note": (
            "Semantic matches are candidate evidence that a concept may be "
            "mentioned. The threshold must be calibrated using labelled "
            "examples. Mention does not prove adequate coverage, student "
            "understanding, mastery, or a prerequisite relationship."
        ),
    }

# =========================================================
# TRANSCRIBE RECORDING AND GENERATE GRAPH TIMELINE
# =========================================================

@app.post("/api/lectures/transcribe")
async def transcribe_lecture(
    file: UploadFile = File(...),
    document_file: UploadFile = File(...),
    language: str = Query("en", min_length=2, max_length=10),
):
    """
    Transcribe an audio/video recording and generate a graph timeline.

    file:
        MP4, MP3, WAV, M4A, WEBM, MPEG, OGG, or FLAC recording

    document_file:
        PDF or PPTX lecture document

    language:
        Speech recognition language code
    """
    try:
        recording_content = await read_upload(
            file,
            MAX_RECORDING_SIZE,
        )
        recording_extension = validate_recording(file.filename)

        document_content = await read_upload(
            document_file,
            MAX_DOCUMENT_SIZE,
        )

        pages = extract_document_pages(
            document_file.filename or "",
            document_content,
        )
        document_text = combine_page_text(pages)

        if not document_text:
            raise HTTPException(
                status_code=422,
                detail=(
                    "No extractable text was found in the lecture document."
                ),
            )

        recording_id = str(uuid4())

        recording_path = (
            RECORDING_DIR
            / f"{recording_id}{recording_extension}"
        )
        recording_path.write_bytes(recording_content)

        # The transcription service must accept:
        # transcribe(path, language)
        raw_result = await run_in_threadpool(
            transcribe,
            recording_path,
            language,
        )

        transcript_data = normalize_transcription_result(raw_result)

        if not transcript_data["segments"]:
            raise HTTPException(
                status_code=422,
                detail=(
                    "Speech recognition returned no segments. "
                    "Check the recording and transcription logs."
                ),
            )

        concepts = extract_concepts(document_text)
        graph = build_knowledge_graph(pages, concepts)

        graph_timeline = build_timeline_compatibility(
            transcript_data,
            pages,
        )

        transcript_path = (
            PROCESSED_DIR
            / f"{recording_id}_transcript.json"
        )

        response_data = {
            "recording_id": recording_id,
            "filename": file.filename,
            "document_filename": document_file.filename,
            "transcript_path": str(transcript_path.resolve()),
            "language": transcript_data["language"] or language,
            "segment_count": len(transcript_data["segments"]),
            "transcript": transcript_data["text"],
            "segments": transcript_data["segments"],
            "concept_count": len(concepts),
            "concepts": concepts,
            "graph": graph,
            "graph_timeline": graph_timeline,
            "message": (
                "Recording transcribed and timestamped graph "
                "timeline generated successfully."
            ),
        }

        transcript_path.write_text(
            json.dumps(
                response_data,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        return response_data

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception(
            "Transcription or graph timeline generation failed"
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "Could not transcribe the recording or generate its graph. "
                "Check the backend terminal for the detailed error."
            ),
        ) from exc
    finally:
        await file.close()
        await document_file.close()
