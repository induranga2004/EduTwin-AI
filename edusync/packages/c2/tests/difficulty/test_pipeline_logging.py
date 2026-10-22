import re

import pytest
from pptx import Presentation

from edusync_c2.difficulty import DifficultyPipeline
from edusync_c2.difficulty.materials import MaterialsError
from edusync_c2.difficulty.schemas import C1GraphResponse
from edusync_c2.difficulty.synthetic import build_c1_graph, build_pptx, make_deck_spec

SPEC = make_deck_spec(3, 6, "lec-log")
FILENAME = "secret_lecture_name.pptx"


@pytest.fixture(scope="module")
def pptx() -> bytes:
    return build_pptx(SPEC)


def graph(version="v1") -> C1GraphResponse:
    return C1GraphResponse.model_validate(build_c1_graph(SPEC, version))


def messages(caplog, level=None):
    return [
        r.getMessage()
        for r in caplog.records
        if r.name.startswith("edusync_c2") and (level is None or r.levelname == level)
    ]


def test_every_step_logs(caplog, pptx):
    caplog.set_level("INFO", logger="edusync_c2")
    DifficultyPipeline().run(pptx, FILENAME, graph())
    recs = [r for r in caplog.records if r.name == "edusync_c2.difficulty.pipeline"]
    msgs = [r.getMessage() for r in recs]
    assert [m.split()[0] for m in msgs] == [
        "materials",
        "graph",
        "concepts",
        "features",
        "fusion",
        "model",
        "done",
    ]
    assert re.fullmatch(r"materials type=pptx bytes=\d+ slides=\d+ warnings=\d+", msgs[0])
    assert re.fullmatch(
        r"graph source=real nodes=6 edges=\d+ prerequisite_edges=\d+ ignored=0", msgs[1]
    )
    assert re.fullmatch(r"concepts total=6 aligned=\d+ not_found=\d+ filter=\w+ dropped=0", msgs[2])
    assert re.fullmatch(r"features concepts=6 computed=\d+ mocked=\d+ unavailable=\d+", msgs[3])
    assert msgs[4] == "fusion inputs=6 features_per_input=26"
    assert msgs[5] == "model name=rule_baseline kind=baseline predicted=6 none=0"
    assert re.fullmatch(r"done concepts=6 duration_ms=\d+", msgs[6])
    assert all(r.levelname == "INFO" for r in recs)


def test_graph_warnings_logged_without_ids_at_warning(caplog, pptx):
    caplog.set_level("DEBUG", logger="edusync_c2")
    raw = build_c1_graph(SPEC, "v1")
    raw["graph"]["edges"].append(
        {
            "source": "ghost-node",
            "target": raw["graph"]["nodes"][0]["id"],
            "relation_type": "prerequisite",
        }
    )
    DifficultyPipeline().run(pptx, FILENAME, C1GraphResponse.model_validate(raw))
    assert "graph_warning count=1" in messages(caplog, "WARNING")
    assert not any("ghost-node" in m for m in messages(caplog, "WARNING"))
    assert any("ghost-node" in m for m in messages(caplog, "DEBUG"))


def test_rejected_input_logged(caplog):
    caplog.set_level("INFO", logger="edusync_c2")
    with pytest.raises(MaterialsError):
        DifficultyPipeline().run(b"hello", FILENAME, None)
    warns = messages(caplog, "WARNING")
    assert len(warns) == 1
    assert warns[0].startswith("materials_rejected status=400")
    assert FILENAME not in warns[0]


def test_privacy_nothing_sensitive_logged(caplog, pptx, tmp_path):
    caplog.set_level("DEBUG", logger="edusync_c2")
    DifficultyPipeline().run(pptx, FILENAME, graph())
    path = tmp_path / "d.pptx"
    path.write_bytes(pptx)
    texts = {
        p.text.strip()
        for slide in Presentation(str(path)).slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
        if len(p.text.strip()) > 25
    }
    assert texts
    assert any(r.levelname == "DEBUG" for r in caplog.records)
    for m in messages(caplog):
        assert FILENAME not in m
        assert "secret_lecture" not in m
        assert not any(t in m for t in texts)


def test_response_independent_of_log_level(caplog, pptx):
    caplog.set_level("DEBUG", logger="edusync_c2")
    debug = DifficultyPipeline().run(pptx, FILENAME, graph()).model_dump_json()
    caplog.set_level("WARNING", logger="edusync_c2")
    quiet = DifficultyPipeline().run(pptx, FILENAME, graph()).model_dump_json()
    assert debug == quiet
    assert "duration" not in debug


def test_info_fields():
    info = DifficultyPipeline(selected_model="minilm").info()
    assert info.selected_model == "minilm"
    assert len(info.feature_schema_hash) == 64
    assert DifficultyPipeline().info().selected_model == "baseline"


def test_logger_is_namespaced():
    from edusync_c2.difficulty import pipeline

    assert pipeline.logger.name == "edusync_c2.difficulty.pipeline"


DISTINCTIVE = (
    "appears",
    "construct",
    "questions",
    "remain",
    "synthetic",
    "testing",
    "stable",
)


@pytest.mark.usefixtures("fresh_default_analyzer")
def test_privacy_vocabulary_analysis_logs_no_words(caplog, pptx, tmp_path):
    caplog.set_level("DEBUG", logger="edusync_c2")
    DifficultyPipeline().run(pptx, FILENAME, graph())
    path = tmp_path / "d.pptx"
    path.write_bytes(pptx)
    slide_text = " ".join(
        p.text.lower()
        for slide in Presentation(str(path)).slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
    )
    words = [w for w in DISTINCTIVE if w in slide_text]
    assert len(words) >= 4  # words that are not concept names (ids are DEBUG by design)
    msgs = messages(caplog)
    assert any(m.startswith("vocab_loaded ") for m in msgs)
    for m in msgs:
        assert not any(w in m.lower() for w in words)


@pytest.mark.usefixtures("fresh_default_analyzer")
def test_vocab_load_logs_once_across_runs(caplog, pptx):
    caplog.set_level("INFO", logger="edusync_c2")
    pipe = DifficultyPipeline()
    pipe.run(pptx, FILENAME, graph())
    pipe.run(pptx, FILENAME, graph())
    assert sum(m.startswith("vocab_loaded ") for m in messages(caplog)) == 1


def test_injected_vocab_reaches_the_features(pptx):
    from edusync_c2.difficulty.features.vocab import VocabStats

    class Fixed:
        def analyze(self, text):
            return VocabStats(words=3, fragmentation=2.5, unfamiliar_share=0.25, domain_share=0.5)

    res = DifficultyPipeline(vocab=Fixed()).run(pptx, FILENAME, graph())  # type: ignore[arg-type]
    for p in res.predictions:
        vals = {f.name: f.value for f in p.features if f.name.startswith("lang.")}
        if vals["lang.avg_word_length"] is None:
            continue
        assert vals["lang.subword_fragmentation"] == 2.5
        assert vals["lang.unfamiliar_vocab_ratio"] == 0.25
        assert vals["lang.domain_term_ratio"] == 0.5
    assert any(p.features[6].value == 2.5 for p in res.predictions)
