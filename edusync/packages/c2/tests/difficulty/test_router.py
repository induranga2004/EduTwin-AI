import importlib
import json
import logging
import re
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from edusync_c2 import COMPONENT
from edusync_c2.difficulty import DifficultyPipeline
from edusync_c2.difficulty.registry import FEATURE_SCHEMA_VERSION
from edusync_c2.difficulty.synthetic import build_c1_graph, build_pdf, build_pptx, make_deck_spec
from edusync_c2.router import get_pipeline
from edusync_core.settings import CoreSettings, get_settings

RT = importlib.import_module("edusync_c2.router")  # "router" on the package is the APIRouter
SPEC = make_deck_spec(5, 5, "lec-r")
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture
def app():
    a = FastAPI()
    a.include_router(COMPONENT.router)
    a.dependency_overrides[get_pipeline] = lambda: DifficultyPipeline()
    return a


@pytest.fixture
def client(app):
    return TestClient(app)


def post(client, data=None, name="deck.pptx", graph=None):
    files = {"file": (name, build_pptx(SPEC) if data is None else data, PPTX_MIME)}
    form = {} if graph is None else {"c1_graph": graph}
    return client.post("/difficulty", files=files, data=form)


def test_with_graph(client):
    r = post(client, graph=json.dumps(build_c1_graph(SPEC, "v1")))
    assert r.status_code == 200
    body = r.json()
    assert body["lecture_id"] == "lec-r"
    assert len(body["predictions"]) == 5
    assert body["inputs"][1]["kind"] == "real"


def test_without_graph(client):
    r = post(client)
    assert r.status_code == 200
    assert r.json()["inputs"][1]["kind"] == "mock"
    assert post(client, graph="  ").json()["inputs"][1]["kind"] == "mock"


def test_pdf(client):
    r = post(client, build_pdf(SPEC), "deck.pdf")
    assert r.status_code == 200
    assert r.json()["file_type"] == "pdf"


def test_wrong_type(client):
    assert post(client, b"hello", "a.txt").status_code == 400


def test_too_large(app, client):
    app.dependency_overrides[get_settings] = lambda: CoreSettings(max_upload_mb=1)
    r = post(client, b"PK" + b"0" * (1024 * 1024 + 10))
    assert r.status_code == 413


def test_invalid_json(client):
    r = post(client, graph="{not json")
    assert r.status_code == 422
    assert "not valid JSON" in r.json()["detail"]


def test_invalid_schema(client):
    r = post(client, graph=json.dumps({"lecture_id": "x"}))
    assert r.status_code == 422
    assert "C1's response shape" in r.json()["detail"]


def test_textless_file(client):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    buf = BytesIO()
    prs.save(buf)
    assert post(client, buf.getvalue()).status_code == 422


def test_old_text_body_rejected(client):
    assert client.post("/difficulty", json={"text": "Hello world."}).status_code == 422


def test_info(client):
    r = client.get("/difficulty/info")
    assert r.status_code == 200
    assert len(r.json()["features"]) == 26
    assert r.json()["model"]["kind"] == "baseline"


def test_health_and_comprehension_untouched(client):
    assert client.get("/health").json() == {"component": "c2", "status": "ok"}
    body = {
        "concept": {"id": "c1", "label": "R"},
        "correct": 1,
        "total": 10,
        "timestamp": {"meeting_id": "m", "seconds": 5},
    }
    assert client.post("/comprehension/evaluate", json=body).json()["flag"]["severity"] == "urgent"


def test_default_pipeline_dependency():
    get_pipeline.cache_clear()
    assert get_pipeline() is get_pipeline()
    assert isinstance(get_pipeline(), DifficultyPipeline)


# ------------------------------------------------------------------ settings, 501s and logs


@pytest.fixture
def real_app(monkeypatch):
    """App using the real (settings-driven) get_pipeline dependency."""
    from edusync_c2.settings import get_c2_settings

    get_c2_settings.cache_clear()
    get_pipeline.cache_clear()
    a = FastAPI()
    a.include_router(COMPONENT.router)
    yield a
    get_c2_settings.cache_clear()
    get_pipeline.cache_clear()


def test_info_shows_model_fields(client):
    body = client.get("/difficulty/info").json()
    assert body["selected_model"] == "baseline"
    assert len(body["feature_schema_hash"]) == 64
    assert body["feature_schema_version"] == FEATURE_SCHEMA_VERSION
    assert {"maturity", "requires", "phase1"} <= set(body["features"][0])


def test_real_dependency_baseline(real_app, monkeypatch):
    monkeypatch.setenv("EDUSYNC_C2_MODEL", "baseline")
    monkeypatch.setattr(RT, "configure_logging", lambda level: None)
    c = TestClient(real_app)
    assert c.get("/difficulty/info").json()["selected_model"] == "baseline"
    assert post(c).status_code == 200


def test_minilm_is_501_and_not_cached(real_app, monkeypatch):
    monkeypatch.setenv("EDUSYNC_C2_MODEL", "minilm")
    monkeypatch.setattr(RT, "configure_logging", lambda level: None)
    c = TestClient(real_app)
    assert c.get("/difficulty/info").status_code == 501
    assert post(c).status_code == 501
    monkeypatch.setenv("EDUSYNC_C2_MODEL", "baseline")
    from edusync_c2.settings import get_c2_settings

    get_c2_settings.cache_clear()
    assert c.get("/difficulty/info").status_code == 200  # the failure was not cached


def test_artifact_error_is_501(real_app, monkeypatch):
    from edusync_c2.difficulty.errors import ModelArtifactError

    def boom(name, model_dir):
        raise ModelArtifactError("artifact folder is broken")

    monkeypatch.setattr(RT, "create_model", boom)
    monkeypatch.setattr(RT, "configure_logging", lambda level: None)
    r = TestClient(real_app).get("/difficulty/info")
    assert (r.status_code, r.json()["detail"]) == (501, "artifact folder is broken")


def test_model_not_implemented_at_inference_is_501(app, client):
    from edusync_c2.difficulty.models import RuleBaselineModel

    class NoInference(RuleBaselineModel):
        def predict(self, inputs):
            raise NotImplementedError("inference is not built yet")

    app.dependency_overrides[get_pipeline] = lambda: DifficultyPipeline(model=NoInference())
    r = post(client)
    assert (r.status_code, r.json()["detail"]) == (501, "inference is not built yet")


def test_configure_logging_called_with_setting_level(real_app, monkeypatch):
    seen = []
    monkeypatch.setenv("EDUSYNC_C2_LOG_LEVEL", "warning")
    monkeypatch.setattr(RT, "configure_logging", seen.append)
    TestClient(real_app).get("/difficulty/info")
    assert seen == ["WARNING"]


def test_request_id_in_log_records_and_rejections_logged(app, client):
    from edusync_c2.logger import RequestIdFilter

    class Grab(logging.Handler):
        def __init__(self):
            super().__init__(logging.DEBUG)
            self.addFilter(RequestIdFilter())
            self.records = []

        def emit(self, record):
            self.records.append(record)

    grab = Grab()
    log = logging.getLogger("edusync_c2")
    old = log.level
    log.addHandler(grab)
    log.setLevel(logging.INFO)
    try:
        assert post(client).status_code == 200
        ok_ids = {r.request_id for r in grab.records}
        assert len(ok_ids) == 1
        assert re.fullmatch(r"[0-9a-f]{8}", ok_ids.pop())
        grab.records.clear()
        assert post(client, b"hello", "secret.txt").status_code == 400
        assert post(client, graph="{nope").status_code == 422
        assert post(client, graph=json.dumps({"lecture_id": "x"})).status_code == 422
        app.dependency_overrides[get_settings] = lambda: CoreSettings(max_upload_mb=1)
        assert post(client, b"PK" + b"0" * (1024 * 1024 + 10)).status_code == 413
        msgs = [r.getMessage() for r in grab.records if r.levelno == logging.WARNING]
        assert any(m.startswith("rejected status=400") for m in msgs)
        assert "rejected status=422 reason=c1_graph_invalid_json" in msgs
        assert "rejected status=422 reason=c1_graph_shape" in msgs
        assert "rejected status=413 reason=too_large limit_mb=1" in msgs
        assert not any("secret" in r.getMessage() for r in grab.records)
    finally:
        log.removeHandler(grab)
        log.setLevel(old)
    assert re.fullmatch(r"-", __import__("edusync_c2.logger", fromlist=["x"]).request_id.get())
