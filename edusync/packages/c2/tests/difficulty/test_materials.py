"""Tests for the PPTX/PDF materials parser and the formula/code detectors."""

from io import BytesIO

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt
from pypdf import PdfWriter

from edusync_c2.difficulty.materials import (
    MaterialsError,
    detect_code,
    detect_formula,
    parse_materials,
)
from edusync_c2.difficulty.synthetic import _png, build_text_pdf


def _save(prs: Presentation) -> bytes:
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
def rich_pptx() -> bytes:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Intro"
    box = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(2))
    box.text_frame.text = "Top line"
    nested = box.text_frame.add_paragraph()
    nested.text = "Nested line"
    nested.level = 2
    box.text_frame.add_paragraph()  # empty paragraph is skipped
    brk = box.text_frame.add_paragraph()
    brk.text = "first"
    brk.add_line_break()
    brk.add_run().text = "second"
    s1.shapes.add_picture(BytesIO(_png()), Inches(6), Inches(2), Inches(1))
    table = s1.shapes.add_table(3, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "a"
    table.cell(0, 1).text = "b"
    table.cell(2, 0).text = "c"  # row 1 stays blank and is skipped
    group = s1.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1)).text_frame.text = "grouped"
    group.shapes.add_picture(BytesIO(_png()), Inches(4), Inches(6), Inches(1))
    s1.notes_slide.notes_text_frame.text = "speaker notes"

    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # no title placeholder
    data = CategoryChartData()
    data.categories = ["x", "y"]
    data.add_series("s", (1, 2))
    s2.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data
    )
    s2.shapes.add_textbox(
        Inches(1), Inches(5), Inches(4), Inches(1)
    ).text_frame.text = "E = m * (c^2)"

    prs.slides.add_slide(prs.slide_layouts[6])  # empty slide
    return _save(prs)


def test_pptx_structure(rich_pptx: bytes) -> None:
    doc = parse_materials(rich_pptx, "Deck.PPTX")
    assert doc.file_type == "pptx"
    assert doc.structure_available
    assert doc.slide_count == len(doc.slides) == 3
    first, second, third = doc.slides
    assert first.index == 1
    assert first.title == "Intro"
    assert first.text.splitlines()[0] == "Intro"
    assert "a | b" in first.text
    assert "c |" in first.text
    assert "grouped" in first.text
    assert "first second" in first.text
    assert first.notes == "speaker notes"
    assert (first.picture_count, first.chart_count, first.table_count) == (2, 0, 1)
    # title, textbox, picture, table, group + its 2 children
    assert first.shape_count == 7
    levels = {p.text: p.level for p in first.paragraphs}
    assert levels["Top line"] == 0
    assert levels["Nested line"] == 2
    assert all(p.text.strip() for p in first.paragraphs)
    assert second.title is None
    assert second.chart_count == 1
    assert second.formula_hint
    assert not second.code_hint
    assert third.text == ""
    assert third.shape_count == 0
    assert doc.warnings == ["Slides without text: 3"]


def test_pdf_pages_and_blank_page_warning() -> None:
    data = build_text_pdf([["Title one", "E = m * (c^2)"], [], ["def f(x):", "    return x;"]])
    doc = parse_materials(data, "notes.pdf")
    assert doc.file_type == "pdf"
    assert not doc.structure_available
    assert doc.slide_count == len(doc.slides) == 3
    first, blank, third = doc.slides
    assert first.title is None
    assert first.picture_count is None
    assert first.shape_count is None
    assert [p.text for p in first.paragraphs] == ["Title one", "E = m * (c^2)"]
    assert all(p.level is None for p in first.paragraphs)
    assert first.formula_hint
    assert third.code_hint
    assert blank.text == ""
    assert blank.paragraphs == []
    assert doc.warnings[0].startswith("PDF: structure fields")
    assert doc.warnings[1] == "PDF pages without extractable text: 2"


def _blank_pdf(encrypt: bool = False) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    if encrypt:
        writer.encrypt("secret")
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_pdf_without_text_is_422() -> None:
    with pytest.raises(MaterialsError) as err:
        parse_materials(_blank_pdf(), "x.pdf")
    assert err.value.status_code == 422


def test_pptx_without_text_is_422() -> None:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    with pytest.raises(MaterialsError) as err:
        parse_materials(_save(prs), "x.pptx")
    assert err.value.status_code == 422


@pytest.mark.parametrize(
    ("data", "name", "fragment"),
    [
        (b"hello", "x.txt", "Unsupported"),
        (b"", "x.pdf", "empty"),
        (b"PK\x03\x04", "x.pdf", "not a PDF"),
        (b"%PDF-1.4", "x.pptx", "not a PPTX"),
        (b"%PDF-1.4 garbage", "x.pdf", "Could not read the PDF"),
        (b"PK\x03\x04 garbage", "x.pptx", "Could not read the PPTX"),
    ],
)
def test_validation_errors(data: bytes, name: str, fragment: str) -> None:
    with pytest.raises(MaterialsError) as err:
        parse_materials(data, name)
    assert err.value.status_code == 400
    assert fragment in err.value.message
    assert str(err.value) == err.value.message


def test_encrypted_pdf_rejected() -> None:
    with pytest.raises(MaterialsError) as err:
        parse_materials(_blank_pdf(encrypt=True), "x.pdf")
    assert err.value.status_code == 400
    assert "Encrypted" in err.value.message


def test_too_large_is_413() -> None:
    with pytest.raises(MaterialsError) as err:
        parse_materials(b"%PDF-" + b"0" * (2 * 1024 * 1024), "x.pdf", max_mb=1)
    assert err.value.status_code == 413


@pytest.mark.parametrize(
    "text",
    [
        "x = \\frac{a}{b}",
        "$$ a + b $$",
        "\\sum_i x_i",
        "total ∑ values",
        "a ≤ b",
        "value >= 3",
        "L(w) = sum((y - w*x)^2) / n",
        "y = a*x^2 + b*x + c",
        "z = (x - mu) / sigma",
        "area = 3 - 1",
    ],
)
def test_detect_formula_true(text: str) -> None:
    assert detect_formula(text)


@pytest.mark.parametrize(
    "text",
    [
        "",
        "Gradient descent minimises the loss function.",
        "x = 5",
        "The result = good (usually) when the data is clean",
        "total = a + b;",
        "if a == b then stop",
        "x = " + "1 + " * 30 + "1",
        "-----",
    ],
)
def test_detect_formula_false(text: str) -> None:
    assert not detect_formula(text)


@pytest.mark.parametrize(
    "text",
    [
        "def run(x):\n    return x;",
        "import os\nvalue = {1: 2}",
        "function f() {\n  return 1;\n}",
        "SELECT name FROM users\nwhere a == b",
        "x -> y;\n\tprint(x)",
    ],
)
def test_detect_code_true(text: str) -> None:
    assert detect_code(text)


@pytest.mark.parametrize(
    "text",
    [
        "",
        "Plain prose about classes and functions.",
        "Please return the form",
        "SELECT only one signal here",
        "y = a*x^2 + b*x + c",
    ],
)
def test_detect_code_false(text: str) -> None:
    assert not detect_code(text)


# ------------------------------------------------------------------ parser regressions
_UNRECOGNISED_SP = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="99" name="Odd shape"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="100" cy="100"/></a:xfrm></p:spPr>'
    "</p:sp>"
)


def test_shape_without_recognisable_geometry_does_not_fail_the_deck() -> None:
    from lxml import etree

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Odd"
    slide.shapes._spTree.append(etree.fromstring(_UNRECOGNISED_SP))
    odd = next(s for s in slide.shapes if s.shape_id == 99)
    with pytest.raises(NotImplementedError):  # the python-pptx behaviour that used to give a 400
        _ = odd.shape_type
    doc = parse_materials(_save(prs), "odd.pptx")
    assert doc.slide_count == 1
    assert doc.slides[0].shape_count == 2
    assert doc.slides[0].picture_count == 0


def test_filled_picture_placeholder_counts_as_a_picture() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])  # Picture with Caption
    picture_ph = next(p for p in slide.placeholders if p.placeholder_format.type.name == "PICTURE")
    slide.shapes.title.text = "Pic"
    picture_ph.insert_picture(BytesIO(_png()))
    doc = parse_materials(_save(prs), "pic.pptx")
    assert doc.slides[0].picture_count == 1


def test_unfilled_picture_placeholder_is_not_a_picture() -> None:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[8]).shapes.title.text = "Empty"
    doc = parse_materials(_save(prs), "empty.pptx")
    assert doc.slides[0].picture_count == 0


# ---------------------------------------------------------------------------- slide geometry

EMU = 914400  # one inch
_NS = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
)
_DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _parse_one(prs: Presentation):
    return parse_materials(_save(prs), "deck.pptx").slides[0]


def _caption(slide) -> None:
    """Text far in a corner so a deck is never rejected for lacking text (area is known)."""
    slide.shapes.add_textbox(0, 0, EMU, EMU).text_frame.text = "caption"


def _graphic_frame(slide, uri: str, *, with_xfrm: bool = True, shape_id: int = 900):
    xfrm = '<p:xfrm><a:off x="100" y="200"/><a:ext cx="3000" cy="4000"/></p:xfrm>'
    xml = (
        f"<p:graphicFrame {_NS}><p:nvGraphicFramePr>"
        f'<p:cNvPr id="{shape_id}" name="Frame"/><p:cNvGraphicFramePr/><p:nvPr/>'
        f"</p:nvGraphicFramePr>{xfrm if with_xfrm else ''}"
        f'<a:graphic><a:graphicData uri="{uri}"/></a:graphic></p:graphicFrame>'
    )
    slide.shapes._spTree.append(parse_xml(xml))


def _set_group(group, off, ext, ch_off, ch_ext) -> None:
    xfrm = group._element.grpSpPr.xfrm
    xfrm.off.x, xfrm.off.y = off
    xfrm.ext.cx, xfrm.ext.cy = ext
    xfrm.chOff.x, xfrm.chOff.y = ch_off
    xfrm.chExt.cx, xfrm.chExt.cy = ch_ext


def _strip_xfrm(shape) -> None:
    spPr = shape._element.spPr
    spPr.remove(spPr.xfrm)


def test_slide_size_and_new_fields_for_pptx() -> None:
    prs = Presentation()
    _caption(_blank_slide(prs))
    doc = parse_materials(_save(prs), "deck.pptx")
    assert (doc.slide_width_emu, doc.slide_height_emu) == (9144000, 6858000)
    slide = doc.slides[0]
    assert slide.text_area_emu == EMU * EMU
    assert slide.picture_area_emu == 0
    assert slide.visual_area_emu == 0
    assert slide.min_font_pt is None
    assert slide.diagram_element_count == 0
    assert slide.table_cell_count == 0
    assert slide.smartart_count == 0
    assert slide.geometry_unknown_count == 0


def test_pdf_leaves_every_geometry_field_none() -> None:
    doc = parse_materials(build_text_pdf([["Title one", "body"]]), "deck.pdf")
    assert doc.slide_width_emu is None
    assert doc.slide_height_emu is None
    slide = doc.slides[0]
    for name in (
        "text_area_emu",
        "picture_area_emu",
        "visual_area_emu",
        "min_font_pt",
        "diagram_element_count",
        "table_cell_count",
        "smartart_count",
        "geometry_unknown_count",
    ):
        assert getattr(slide, name) is None, name


def test_textbox_rect_and_font_sizes() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(EMU, EMU, 4 * EMU, 2 * EMU)
    box.text_frame.text = "run sized"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    second = box.text_frame.add_paragraph()
    second.text = "paragraph sized"
    second.font.size = Pt(9)  # paragraph level only; the run carries no size
    third = box.text_frame.add_paragraph()
    third.text = "run wins"
    third.font.size = Pt(5)
    third.runs[0].font.size = Pt(20)
    blank_run = box.text_frame.add_paragraph()
    blank_run.add_run().text = "  "
    blank_run.runs[0].font.size = Pt(3)  # whitespace-only runs are ignored
    record = _parse_one(prs)
    assert record.min_font_pt == 9.0
    assert record.text_area_emu == 4 * EMU * 2 * EMU
    assert record.geometry_unknown_count == 0


def test_run_size_beats_a_smaller_paragraph_size() -> None:
    prs = Presentation()
    box = _blank_slide(prs).shapes.add_textbox(0, 0, EMU, EMU)
    box.text_frame.text = "x"
    box.text_frame.paragraphs[0].font.size = Pt(5)
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    assert _parse_one(prs).min_font_pt == 20.0


def test_min_font_is_smallest_across_shapes_and_none_without_explicit_sizes() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    for size in (Pt(14), Pt(10.5), Pt(18)):
        box = slide.shapes.add_textbox(0, 0, EMU, EMU)
        box.text_frame.text = "t"
        box.text_frame.paragraphs[0].runs[0].font.size = size
    assert _parse_one(prs).min_font_pt == 10.5
    plain = Presentation()
    _blank_slide(plain).shapes.add_textbox(0, 0, EMU, EMU).text_frame.text = "no size"
    assert _parse_one(plain).min_font_pt is None


def test_picture_chart_and_text_areas_are_separate() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    picture = slide.shapes.add_picture(BytesIO(_png()), 2 * EMU, 3 * EMU, EMU)
    data = CategoryChartData()
    data.categories = ["x", "y"]
    data.add_series("s", (1, 2))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, 3 * EMU, 2 * EMU, data)
    _caption(slide)
    record = _parse_one(prs)
    area = picture.width * picture.height
    assert record.picture_area_emu == area
    assert record.visual_area_emu == area  # the chart is neither a picture nor a diagram leaf
    assert record.text_area_emu == EMU * EMU  # the picture is not text
    assert record.diagram_element_count == 0
    assert (record.picture_count, record.chart_count) == (1, 1)


def test_table_cells_include_merged_cells_and_the_frame_is_a_text_rect() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    frame = slide.shapes.add_table(3, 3, EMU, EMU, 3 * EMU, 3 * EMU)
    table = frame.table
    table.cell(0, 0).text = "head"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 1).text_frame.text = "small"
    table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(7)
    empty = slide.shapes.add_table(2, 2, 5 * EMU, 5 * EMU, EMU, EMU)  # no text anywhere
    assert empty.table is not None
    record = _parse_one(prs)
    assert record.table_cell_count == 9 + 4
    assert record.table_count == 2
    assert record.text_area_emu == frame.width * frame.height
    assert record.min_font_pt == 7.0
    assert record.diagram_element_count == 0


def test_connector_autoshape_and_freeform_are_diagram_leaves() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, EMU, EMU, 2 * EMU, 2 * EMU)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 3 * EMU, EMU, EMU, EMU)
    builder = slide.shapes.build_freeform(5 * EMU, EMU)
    builder.add_line_segments([(6 * EMU, EMU), (6 * EMU, 2 * EMU)])
    freeform = builder.convert_to_shape()
    slide.shapes.add_textbox(0, 4 * EMU, EMU, EMU).text_frame.text = "text box"  # not a leaf
    record = _parse_one(prs)
    assert record.diagram_element_count == 3
    expected = EMU * EMU + EMU * EMU + freeform.width * freeform.height
    assert record.visual_area_emu == expected
    assert record.picture_area_emu == 0
    assert record.smartart_count == 0
    assert record.geometry_unknown_count == 0


def test_text_inside_an_autoshape_is_text_and_diagram() -> None:
    prs = Presentation()
    shape = _blank_slide(prs).shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 2 * EMU, EMU)
    shape.text_frame.text = "node"
    record = _parse_one(prs)
    assert record.text_area_emu == record.visual_area_emu == 2 * EMU * EMU
    assert record.diagram_element_count == 1


def test_smartart_frame_counts_as_a_diagram_leaf_with_its_rect() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    _caption(slide)
    _graphic_frame(slide, _DIAGRAM_URI)
    _graphic_frame(slide, "http://example.com/other", shape_id=901)  # unrelated graphic data
    record = _parse_one(prs)
    assert record.smartart_count == 1
    assert record.diagram_element_count == 1
    assert record.visual_area_emu == 3000 * 4000
    assert record.picture_area_emu == 0
    assert record.geometry_unknown_count == 0


def test_graphic_frame_without_xfrm_is_unknown_but_still_counted() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    _caption(slide)
    _graphic_frame(slide, _DIAGRAM_URI, with_xfrm=False)
    record = _parse_one(prs)
    assert record.smartart_count == 1
    assert record.diagram_element_count == 1
    assert record.geometry_unknown_count == 1
    assert record.visual_area_emu == 0


def test_shape_with_stripped_xfrm_is_unknown() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    _caption(slide)
    lost = slide.shapes.add_textbox(0, 0, EMU, EMU)
    lost.text_frame.text = "lost"
    _strip_xfrm(lost)
    empty = slide.shapes.add_textbox(0, 0, EMU, EMU)  # no text: its rect is never needed
    _strip_xfrm(empty)
    record = _parse_one(prs)
    assert record.geometry_unknown_count == 1
    assert record.text_area_emu == EMU * EMU
    assert record.shape_count == 3


def test_group_children_use_the_group_transform() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 1000, 1000, 2000, 2000)
    _set_group(group, (10000, 20000), (4000, 6000), (1000, 1000), (2000, 2000))
    _caption(slide)
    record = _parse_one(prs)
    assert record.visual_area_emu == 4000 * 6000
    assert record.diagram_element_count == 1  # the group itself is not counted
    assert record.shape_count == 3  # ... but it is still a shape, as before


def test_nested_rotated_shape_in_nested_groups() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    node = inner.shapes.add_shape(MSO_SHAPE.RECTANGLE, 1000, 1000, 2000, 1000)
    node.rotation = 90
    _set_group(inner, (10000, 10000), (4000, 3000), (1000, 1000), (2000, 1000))
    _set_group(outer, (1000, 2000), (100000, 100000), (0, 0), (100000, 100000))
    _caption(slide)
    record = _parse_one(prs)
    assert record.visual_area_emu == 2000 * 6000
    assert record.geometry_unknown_count == 0


def test_group_with_zero_child_extent_makes_children_unknown() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 1000, 1000, 2000, 2000)
    _set_group(group, (0, 0), (4000, 6000), (1000, 1000), (0, 0))
    _caption(slide)
    record = _parse_one(prs)
    assert record.geometry_unknown_count == 1
    assert record.diagram_element_count == 1
    assert record.visual_area_emu == 0


def test_group_without_transform_makes_children_unknown() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, EMU, EMU).text_frame.text = "inside"
    group.shapes.add_textbox(0, 0, EMU, EMU)  # no text: its rect is never needed
    group._element.grpSpPr.remove(group._element.grpSpPr.xfrm)
    record = _parse_one(prs)
    assert record.geometry_unknown_count == 1
    assert record.text_area_emu == 0


def test_placeholders_use_inherited_geometry_and_are_not_diagram_leaves() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Body"
    record = _parse_one(prs)
    title, body = slide.shapes.title, slide.placeholders[1]
    assert record.diagram_element_count == 0
    assert record.geometry_unknown_count == 0
    assert 0 < record.text_area_emu <= max(title.width * title.height + body.width * body.height, 1)
    assert record.shape_count == 2


def test_placeholder_with_broken_inheritance_is_unknown() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Body"
    for holder in (*prs.slide_layouts[1].placeholders, *prs.slide_master.placeholders):
        spPr = holder._element.spPr
        if spPr.xfrm is not None:
            spPr.remove(spPr.xfrm)
    record = _parse_one(prs)
    assert record.geometry_unknown_count == 2  # the title and the body
    assert record.text_area_emu == 0
    assert record.title == "Title"


def test_title_font_and_rect_are_collected_before_the_title_is_skipped() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Big title"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(6)
    record = _parse_one(prs)
    assert record.min_font_pt == 6.0
    assert record.text_area_emu > 0
    assert record.shape_count == 1


@pytest.mark.parametrize("size", [None, 0])
def test_unknown_or_zero_slide_size_leaves_the_areas_none(size: int | None) -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    slide.shapes.add_textbox(0, 0, EMU, EMU).text_frame.text = "x"
    slide.shapes.add_table(2, 2, 0, 0, EMU, EMU)
    if size is None:
        prs._element.remove(prs._element.sldSz)
    else:
        prs._element.sldSz.set("cx", str(size))
    doc = parse_materials(_save(prs), "deck.pptx")
    record = doc.slides[0]
    assert record.text_area_emu is None
    assert record.picture_area_emu is None
    assert record.visual_area_emu is None
    assert record.table_cell_count == 4  # counts do not need the slide size
    assert doc.slide_width_emu == size


def test_shapes_off_the_slide_are_clamped() -> None:
    prs = Presentation()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(-EMU, 0, 2 * EMU, EMU)
    box.text_frame.text = "half off"
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 50 * EMU, 0, EMU, EMU)  # entirely off
    record = _parse_one(prs)
    assert record.text_area_emu == EMU * EMU
    assert record.visual_area_emu == 0
    assert record.diagram_element_count == 1
