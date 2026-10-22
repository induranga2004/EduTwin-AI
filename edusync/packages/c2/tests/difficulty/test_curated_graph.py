"""Curated graph provenance (provider, pipeline) and the make_curated_graph draft tool."""

import importlib.util
import json
import runpy
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from edusync_c2.difficulty.pipeline import (
    CURATED_GRAPH_WARNING,
    MOCK_GRAPH_WARNING,
    DifficultyPipeline,
)
from edusync_c2.difficulty.providers import (
    FileGraphProvider,
    curated_graph_path,
    is_curated,
)
from edusync_c2.difficulty.schemas import C1GraphResponse, MaterialsDoc

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "make_curated_graph.py"


@pytest.fixture(scope="module")
def script():
    spec = importlib.util.spec_from_file_location("make_curated_graph", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_deck() -> bytes:
    prs = Presentation()
    rows = [
        ("Overview", ""),
        ("Gradient Descent", "Minimises a loss."),
        ("Learning Rate", "Step size of gradient descent."),
        ("Backpropagation", "Uses the learning rate too."),
        ("gradient   descent", ""),
        ("Questions?", ""),
        ("***", ""),
    ]
    for title, body in rows:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        if body:
            box = slide.shapes.add_textbox(0, 0, 100, 100)
            box.text_frame.text = body
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


DECK = make_deck()


@pytest.fixture
def deck_file(tmp_path: Path) -> Path:
    path = tmp_path / "Lecture 1.pptx"
    path.write_bytes(DECK)
    return path


# ---------------------------------------------------------------------------------- the flag


def _graph(**top) -> C1GraphResponse:
    return C1GraphResponse.model_validate({"graph": {"nodes": [], "edges": []}, **top})


def test_curated_flag_forms() -> None:
    assert is_curated(_graph(curated=True))
    assert is_curated(_graph(note="curated"))
    assert not is_curated(_graph())
    assert not is_curated(_graph(curated=False))
    assert not is_curated(_graph(curated="true"))
    assert not is_curated(_graph(note="Curated by hand"))


def test_curated_graph_path() -> None:
    d = Path("graphs")
    assert curated_graph_path(d, "Lecture 1.pptx") == d / "Lecture 1.graph.json"
    assert curated_graph_path(d, "a\\b\\deck.v2.pdf") == d / "deck.v2.graph.json"
    assert FileGraphProvider.for_deck(d, "x.pptx")._path == d / "x.graph.json"


def _write_graph(tmp_path: Path, deck_file: Path, **top) -> Path:
    draft = _draft(deck_file)
    draft.update(top)
    path = tmp_path / "g.graph.json"
    path.write_text(json.dumps(draft), encoding="utf-8")
    return path


def _draft(deck_file: Path, **kw) -> dict:
    return script_module().build_draft(deck_file.read_bytes(), deck_file.name, **kw)


_MODULE = []


def script_module():
    if not _MODULE:
        spec = importlib.util.spec_from_file_location("make_curated_graph", SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        _MODULE.append(module)
    return _MODULE[0]


# -------------------------------------------------------------------- provider and pipeline


def test_provider_kinds(tmp_path: Path, deck_file: Path, deck: MaterialsDoc) -> None:
    curated = FileGraphProvider(_write_graph(tmp_path, deck_file, curated=True)).build(deck)
    assert curated.source.kind == "curated"
    assert curated.source.detail == "graph file: g.graph.json (curated)"
    by_note = FileGraphProvider(
        _write_graph(tmp_path, deck_file, curated=False, note="curated")
    ).build(deck)
    assert by_note.source.kind == "curated"
    draft = FileGraphProvider(_write_graph(tmp_path, deck_file)).build(deck)
    assert (draft.source.kind, draft.source.detail) == ("mock", "graph file: g.graph.json")
    sample = FileGraphProvider().build(deck)
    assert sample.source.kind == "mock"


def _run(path: Path) -> dict:
    pipeline = DifficultyPipeline(
        graph_provider_factory=lambda data, name: FileGraphProvider(path),
        graph_provider_name="file",
    )
    resp = pipeline.run(DECK, "Lecture 1.pptx")
    assert resp.model_dump() == pipeline.run(DECK, "Lecture 1.pptx").model_dump()
    return resp.model_dump(mode="json")


def _b_and_freq(body: dict) -> set[tuple[str, str]]:
    return {
        (f["name"], f["status"])
        for p in body["predictions"]
        for f in p["features"]
        if f["name"].startswith("prereq.") or f["name"] == "dens.concept_frequency"
    }


def test_curated_graph_through_the_pipeline(tmp_path: Path, deck_file: Path) -> None:
    body = _run(_write_graph(tmp_path, deck_file, curated=True))
    assert body["inputs"][1]["kind"] == "curated"
    assert CURATED_GRAPH_WARNING in body["warnings"]
    assert MOCK_GRAPH_WARNING not in body["warnings"]
    statuses = _b_and_freq(body)
    assert ("prereq.direct_count", "computed") in statuses
    assert ("dens.concept_frequency", "computed") in statuses
    assert not any(s == "mocked" for _, s in statuses)
    graph_stage = next(s for s in body["stages"] if s["name"] == "graph")
    assert graph_stage["status"] == "ok"
    assert "(curated)" in graph_stage["detail"]
    sources = {
        f["source"] for p in body["predictions"] for f in p["features"] if "prereq" in f["name"]
    }
    assert not any(s.startswith("mock:") for s in sources)


def test_uncurated_file_stays_mock(tmp_path: Path, deck_file: Path) -> None:
    body = _run(_write_graph(tmp_path, deck_file))
    assert body["inputs"][1]["kind"] == "mock"
    assert MOCK_GRAPH_WARNING in body["warnings"]
    assert CURATED_GRAPH_WARNING not in body["warnings"]
    assert ("prereq.direct_count", "mocked") in _b_and_freq(body)


# ----------------------------------------------------------------------------- the draft tool


def test_draft_content_and_determinism(deck_file: Path) -> None:
    draft = _draft(deck_file)
    assert draft == _draft(deck_file)
    assert draft["curated"] is False
    assert draft["review"]
    assert draft["lecture_id"] == "lecture-1"
    nodes = draft["graph"]["nodes"]
    # Overview (generic), duplicate title, 'Questions?' and '***' are skipped
    assert [n["id"] for n in nodes] == ["gradient-descent", "learning-rate", "backpropagation"]
    assert [n["frequency"] for n in nodes] == [3, 2, 1]  # slides mentioning the label
    assert draft["concept_count"] == draft["graph"]["node_count"] == 3
    edges = draft["graph"]["edges"]
    assert [(e["source"], e["target"]) for e in edges] == [
        ("gradient-descent", "learning-rate"),
        ("learning-rate", "backpropagation"),
    ]
    assert all(e["confidence"] == 0.5 and e["source_type"] == "inferred" for e in edges)
    assert edges[0]["evidence"][0]["page"] == 3


def _deck(rows: list[tuple[str, str]]) -> bytes:
    prs = Presentation()
    for title, body in rows:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        if body:
            box = slide.shapes.add_textbox(0, 0, 100, 100)
            box.text_frame.text = body
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_draft_mention_edges_confidence_cap_and_roots() -> None:
    names = ["Alpha", "Bravo", "Charlie", "Delta", "Echo", "Foxtrot"]
    rows = [(n, "") for n in names]
    rows.append(("Zulu", "alpha bravo charlie delta echo foxtrot"))
    rows.append(("Zulu", "ALPHA and Bravo again"))
    rows.append(("Lonely", "Nothing earlier is named here"))
    draft = script_module().build_draft(_deck(rows), "d.pptx")
    edges = {(e["source"], e["target"]): e for e in draft["graph"]["edges"]}
    into_zulu = {s: e for (s, t), e in edges.items() if t == "zulu"}
    assert sorted(into_zulu) == ["alpha", "bravo", "charlie", "delta", "echo"]  # capped at 5
    assert into_zulu["alpha"]["confidence"] == into_zulu["bravo"]["confidence"] == 0.6
    assert into_zulu["echo"]["confidence"] == 0.5
    assert into_zulu["alpha"]["evidence"] == [{"page": 7, "sentence": ""}]
    assert all(e["source_type"] == "inferred" for e in edges.values())
    assert not any(t == "lonely" for _, t in edges)  # no fallback chain: a root
    assert all(
        {n["id"]: i for i, n in enumerate(draft["graph"]["nodes"])}[s]
        < {n["id"]: i for i, n in enumerate(draft["graph"]["nodes"])}[t]
        for s, t in edges
    )


def test_draft_confidence_caps_at_point_eight() -> None:
    rows = [("Alpha", ""), ("Zulu", "alpha")] + [("Zulu", "alpha")] * 4
    draft = script_module().build_draft(_deck(rows), "d.pptx")
    (edge,) = draft["graph"]["edges"]
    assert edge["confidence"] == 0.8


def test_draft_min_frequency(deck_file: Path) -> None:
    ids = [n["id"] for n in _draft(deck_file, min_frequency=2)["graph"]["nodes"]]
    assert ids == ["gradient-descent", "learning-rate"]


def test_prereq_map_overrides_and_extends(deck_file: Path) -> None:
    draft = _draft(
        deck_file,
        prereq_map={"Backpropagation": ["Gradient Descent", "Learning Rate"]},
    )
    edges = {(e["source"], e["target"]): e for e in draft["graph"]["edges"]}
    assert len(edges) == 3
    for key in (("learning-rate", "backpropagation"), ("gradient-descent", "backpropagation")):
        assert (edges[key]["confidence"], edges[key]["source_type"]) == (0.9, "curriculum_defined")
    assert edges[("gradient-descent", "learning-rate")]["source_type"] == "inferred"


def test_prereq_map_unknown_label(deck_file: Path) -> None:
    for bad in ({"Nope": ["Learning Rate"]}, {"Learning Rate": ["Nope"]}):
        with pytest.raises(ValueError, match="not a concept"):
            _draft(deck_file, prereq_map=bad)


def test_cli_writes_a_valid_draft_and_refuses_overwrite(
    script, deck_file: Path, tmp_path: Path, capsys
) -> None:
    out = tmp_path / "out" / "Lecture 1.graph.json"
    out.parent.mkdir()
    argv = ["--deck", str(deck_file), "--out", str(out)]
    assert script.main(argv) == 0
    assert "Wrote" in capsys.readouterr().out
    written = json.loads(out.read_text(encoding="utf-8"))
    assert written["curated"] is False
    first = out.read_text(encoding="utf-8")
    assert script.main(argv) == 2
    assert "--force" in capsys.readouterr().err
    assert out.read_text(encoding="utf-8") == first
    assert script.main([*argv, "--force", "--min-frequency", "2"]) == 0
    assert len(json.loads(out.read_text(encoding="utf-8"))["graph"]["nodes"]) == 2


def test_cli_prereq_map_file(script, deck_file: Path, tmp_path: Path) -> None:
    pm = tmp_path / "pm.json"
    pm.write_text(json.dumps({"Backpropagation": ["Gradient Descent"]}), encoding="utf-8")
    out = tmp_path / "o.json"
    assert script.main(["--deck", str(deck_file), "--out", str(out), "--prereq-map", str(pm)]) == 0
    edges = json.loads(out.read_text(encoding="utf-8"))["graph"]["edges"]
    assert any(e["source_type"] == "curriculum_defined" for e in edges)


def test_cli_validator_errors_exit_one(script, deck_file: Path, tmp_path: Path, capsys) -> None:
    pm = tmp_path / "pm.json"
    pm.write_text(json.dumps({"Learning Rate": ["Learning Rate"]}), encoding="utf-8")
    out = tmp_path / "o.json"
    argv = ["--deck", str(deck_file), "--out", str(out), "--prereq-map", str(pm)]
    assert script.main(argv) == 1  # self-edge
    err = capsys.readouterr().err
    assert "C1-008" in err
    assert not out.exists()


def test_script_entry_point(deck_file: Path, tmp_path: Path, monkeypatch) -> None:
    out = tmp_path / "entry.json"
    monkeypatch.setattr(
        "sys.argv", ["make_curated_graph.py", "--deck", str(deck_file), "--out", str(out)]
    )
    with pytest.raises(SystemExit) as exc:
        runpy.run_path(str(SCRIPT), run_name="__main__")
    assert exc.value.code == 0
    assert out.exists()


def test_cli_unreadable_inputs_exit_two(script, deck_file: Path, tmp_path: Path, capsys) -> None:
    out = tmp_path / "o.json"
    assert script.main(["--deck", str(tmp_path / "none.pptx"), "--out", str(out)]) == 2
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a deck")
    assert script.main(["--deck", str(bad), "--out", str(out)]) == 2
    pm = tmp_path / "pm.json"
    pm.write_text("{", encoding="utf-8")
    assert script.main(["--deck", str(deck_file), "--out", str(out), "--prereq-map", str(pm)]) == 2
    pm.write_text(json.dumps({"Nope": ["Learning Rate"]}), encoding="utf-8")
    assert script.main(["--deck", str(deck_file), "--out", str(out), "--prereq-map", str(pm)]) == 2
    assert "Cannot build the draft" in capsys.readouterr().err
    assert not out.exists()
