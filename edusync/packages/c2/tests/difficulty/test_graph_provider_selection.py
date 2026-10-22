"""Settings-driven graph provider selection, end to end through the HTTP route."""

import importlib
import json
import logging
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from edusync_c2 import COMPONENT
from edusync_c2.difficulty.errors import GraphProviderError
from edusync_c2.difficulty.pipeline import MOCK_GRAPH_WARNING, DifficultyPipeline
from edusync_c2.difficulty.providers import C1ApiGraphProvider, FileGraphProvider
from edusync_c2.difficulty.synthetic import build_c1_graph, build_pptx, make_deck_spec
from edusync_c2.router import _graph_provider_factory, get_pipeline
from edusync_c2.settings import C2Settings, get_c2_settings

RT = importlib.import_module("edusync_c2.router")
SPEC = make_deck_spec(5, 5, "lec-sel")
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DECK = build_pptx(SPEC)
V1 = build_c1_graph(SPEC, "v1")


def settings(**kw) -> C2Settings:
    return C2Settings(_env_file=None, **kw)


def post(client, data=DECK, name="secret-name.pptx", graph=None):
    files = {"file": (name, data, PPTX_MIME)}
    form = {} if graph is None else {"c1_graph": json.dumps(graph)}
    return client.post("/difficulty", files=files, data=form)


@pytest.fixture
def real_app(monkeypatch):
    monkeypatch.setattr(RT, "configure_logging", lambda level: None)
    get_c2_settings.cache_clear()
    get_pipeline.cache_clear()
    app = FastAPI()
    app.include_router(COMPONENT.router)
    yield app
    get_c2_settings.cache_clear()
    get_pipeline.cache_clear()


class FakeTransport:
    """Stands in for UrllibTransport; class attributes steer the outcome."""

    outcome: object = V1
    seen: list = []  # noqa: RUF012

    def __init__(self, base_url: str, timeout_s: float) -> None:
        FakeTransport.seen.append(("init", base_url, timeout_s))

    def __call__(self, data: bytes, filename: str) -> dict:
        FakeTransport.seen.append(("call", len(data), filename))
        if isinstance(self.outcome, Exception):
            raise self.outcome
        return self.outcome  # type: ignore[return-value]


@pytest.fixture
def fake(monkeypatch):
    FakeTransport.outcome = V1
    FakeTransport.seen = []
    monkeypatch.setattr(RT, "UrllibTransport", FakeTransport)
    return FakeTransport


# ------------------------------------------------------------------ settings


def test_settings_defaults(monkeypatch):
    for k in ("GRAPH_PROVIDER", "GRAPH_FILE", "C1_BASE_URL", "C1_TIMEOUT_S"):
        monkeypatch.delenv(f"EDUSYNC_C2_{k}", raising=False)
    s = settings()
    assert (s.graph_provider, s.graph_file) == ("mock", None)
    assert (s.c1_base_url, s.c1_timeout_s) == ("http://127.0.0.1:8100/api/c1", 60)


def test_settings_env(monkeypatch, tmp_path):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "c1_api")
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_FILE", str(tmp_path / "g.json"))
    monkeypatch.setenv("EDUSYNC_C2_C1_BASE_URL", " https://c1.example/api/c1 ")
    monkeypatch.setenv("EDUSYNC_C2_C1_TIMEOUT_S", "2.5")
    s = settings()
    assert s.graph_provider == "c1_api"
    assert s.graph_file == tmp_path / "g.json"
    assert (s.c1_base_url, s.c1_timeout_s) == ("https://c1.example/api/c1", 2.5)


def test_empty_graph_file_is_none(monkeypatch):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_FILE", "")
    assert settings().graph_file is None


@pytest.mark.parametrize("url", ["file:///etc/passwd", "ftp://x", "c1.local:8100", ""])
def test_base_url_scheme_rejected(url):
    with pytest.raises(ValueError, match="c1_base_url"):
        settings(c1_base_url=url)


@pytest.mark.parametrize("kw", [{"graph_provider": "ftp"}, {"c1_timeout_s": 0}])
def test_invalid_provider_and_timeout(kw):
    with pytest.raises(ValueError, match=r"graph_provider|c1_timeout_s"):
        settings(**kw)


# ------------------------------------------------------------------ factory


def test_factory_selection(tmp_path):
    assert _graph_provider_factory(settings()) is None
    file_factory = _graph_provider_factory(settings(graph_provider="file"))
    assert file_factory is not None
    provider = file_factory(b"x", "a.pptx")
    assert isinstance(provider, FileGraphProvider)
    api_factory = _graph_provider_factory(settings(graph_provider="c1_api"))
    assert api_factory is not None
    assert isinstance(api_factory(b"x", "a.pptx"), C1ApiGraphProvider)


def test_pipeline_info_reports_configured_name():
    assert DifficultyPipeline().info().graph_provider == "mock_from_slide_titles"
    assert DifficultyPipeline(graph_provider_name="file").info().graph_provider == "file"


def test_info_route_reports_provider(real_app, monkeypatch):
    c = TestClient(real_app)
    assert c.get("/difficulty/info").json()["graph_provider"] == "mock_from_slide_titles"
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "file")
    get_c2_settings.cache_clear()
    get_pipeline.cache_clear()
    assert c.get("/difficulty/info").json()["graph_provider"] == "file"


# ------------------------------------------------------------------ c1_api end to end


def test_c1_api_end_to_end_deterministic(real_app, monkeypatch, fake):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "c1_api")
    monkeypatch.setenv("EDUSYNC_C2_C1_BASE_URL", "http://c1.test/api/c1")
    monkeypatch.setenv("EDUSYNC_C2_C1_TIMEOUT_S", "7")
    c = TestClient(real_app)
    r1, r2 = post(c), post(c)
    assert r1.status_code == r2.status_code == 200
    assert r1.json() == r2.json()
    body = r1.json()
    assert body["lecture_id"].startswith("c1-")
    assert len(body["lecture_id"]) == 15
    graph_input = body["inputs"][1]
    assert (graph_input["name"], graph_input["kind"]) == ("c1_graph", "real")
    assert graph_input["detail"].startswith("C1 API graph")
    assert MOCK_GRAPH_WARNING not in body["warnings"]
    assert fake.seen[0] == ("init", "http://c1.test/api/c1", 7.0)
    assert fake.seen[1] == ("call", len(DECK), "secret-name.pptx")
    groups = {f["name"]: f["status"] for p in body["predictions"] for f in p["features"]}
    assert groups["prereq.direct_count"] == "computed"
    assert groups["prereq.chain_depth"] == "computed"
    assert all(p["difficulty"] is not None for p in body["predictions"])


def test_posted_graph_beats_configured_provider(real_app, monkeypatch, fake):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "c1_api")
    c = TestClient(real_app)
    r = post(c, graph=V1)
    assert r.status_code == 200
    assert r.json()["lecture_id"] == "lec-sel"
    assert fake.seen == []


@pytest.mark.parametrize(
    ("status", "message"),
    [(422, "C1 rejected the file"), (502, "C1 is unreachable"), (504, "C1 timed out")],
)
def test_provider_errors_are_mapped(real_app, monkeypatch, fake, caplog, status, message):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "c1_api")
    fake.outcome = GraphProviderError(status, message)
    caplog.set_level(logging.DEBUG, logger="edusync_c2")
    r = post(TestClient(real_app))
    assert (r.status_code, r.json()["detail"]) == (status, message)
    rejected = [m.getMessage() for m in caplog.records if "rejected" in m.getMessage()]
    assert rejected == [f"rejected status={status} reason=graph_provider"]
    logged = " ".join(m.getMessage() for m in caplog.records)
    assert "secret-name" not in logged


def test_invalid_c1_response_is_502(real_app, monkeypatch, fake):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "c1_api")
    fake.outcome = {"graph": "nope"}
    assert post(TestClient(real_app)).status_code == 502


# ------------------------------------------------------------------ file provider end to end


def test_file_provider_end_to_end(real_app, monkeypatch):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "file")
    c = TestClient(real_app)
    prs = Presentation()
    for title in ("Cryptography basics", "Nonce reuse"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    buf = BytesIO()
    prs.save(buf)
    r1, r2 = post(c, buf.getvalue()), post(c, buf.getvalue())
    assert r1.status_code == 200
    assert r1.json() == r2.json()
    body = r1.json()
    assert body["inputs"][1]["kind"] == "mock"
    assert body["inputs"][1]["detail"] == "graph file: sample_c1_graph_v1.json"
    assert MOCK_GRAPH_WARNING in body["warnings"]
    assert body["lecture_id"]
    flagged = [p for p in body["predictions"] if p["alignment_status"] == "not_found_in_materials"]
    assert flagged
    assert all("concept not found in the lecture materials" in p["warnings"] for p in flagged)


def test_file_provider_missing_file_is_500(real_app, monkeypatch, tmp_path):
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_PROVIDER", "file")
    monkeypatch.setenv("EDUSYNC_C2_GRAPH_FILE", str(tmp_path / "missing.json"))
    r = post(TestClient(real_app))
    assert r.status_code == 500
    assert "missing.json" not in r.text


# ------------------------------------------------------------------ default unchanged


def test_default_is_mock_with_no_lecture_id(real_app):
    c = TestClient(real_app)
    body = post(c).json()
    assert body["lecture_id"] is None
    assert body["inputs"][1]["kind"] == "mock"
