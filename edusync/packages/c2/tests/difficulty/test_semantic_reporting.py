"""Semantic scorer reporting and failure mapping, end to end through the HTTP routes."""

import importlib
import logging
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from edusync_c2 import COMPONENT
from edusync_c2.difficulty import DifficultyPipeline
from edusync_c2.difficulty.errors import SemanticScoringError
from edusync_c2.difficulty.features.semantic import SemanticScorer
from edusync_c2.difficulty.pipeline import (
    SEMANTIC_UNAVAILABLE_WARNING,
)
from edusync_c2.router import get_pipeline
from edusync_c2.settings import get_c2_settings

RT = importlib.import_module("edusync_c2.router")
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
SECRET = "zebra quokka marmalade"
SEMANTIC = (
    "expl.semantic_definition_count",
    "expl.semantic_example_count",
    "expl.semantic_support_score",
)


def build_deck() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Gradient descent"
    slide.placeholders[1].text_frame.text = (
        f"DEFN gradient descent is an optimiser for {SECRET}.\n"
        f"EXMP for instance run it once on {SECRET} data.\n"
        "Some plain filler text goes here."
    )
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


DECK = build_deck()


def post(client, data=DECK):
    return client.post("/difficulty", files={"file": ("deck.pptx", data, PPTX_MIME)})


def feature(body, name):
    return next(f for f in body["predictions"][0]["features"] if f["name"] == name)


@pytest.fixture
def real_app(monkeypatch):
    monkeypatch.setattr(RT, "configure_logging", lambda level: None)
    for key in ("SEMANTIC", "SEMANTIC_MODEL_DIR"):
        monkeypatch.delenv(f"EDUSYNC_C2_{key}", raising=False)
    get_c2_settings.cache_clear()
    get_pipeline.cache_clear()
    app = FastAPI()
    app.include_router(COMPONENT.router)
    yield app
    get_c2_settings.cache_clear()
    get_pipeline.cache_clear()


@pytest.fixture
def minilm(real_app, monkeypatch, keyword_encoder):
    """Settings select minilm and the loader hands out the keyword encoder."""
    monkeypatch.setenv("EDUSYNC_C2_SEMANTIC", "minilm")
    monkeypatch.setattr(RT, "get_default_encoder", lambda model_dir: keyword_encoder)
    return TestClient(real_app)


class PrototypesOnly:
    """Encodes the prototype sentences, then misbehaves on any slide sentence."""

    dims = 4
    sha = "feedbeef"

    def __init__(self, inner, mode: str) -> None:
        self.inner = inner
        self.mode = mode

    def encode(self, texts):
        if self.mode == "ok" or all(t in self.inner.axis for t in texts):
            return self.inner.encode(texts)
        if self.mode == "raise":
            raise RuntimeError(f"boom: {' | '.join(texts)}")
        return self.inner.encode(texts)[:-1]  # one row short


# ------------------------------------------------------------------ /info and warning


def test_info_defaults_to_off(real_app):
    r = TestClient(real_app).get("/difficulty/info")
    assert r.status_code == 200
    assert r.json()["semantic_scorer"] == "off"


def test_info_shows_the_weights_fingerprint(minilm):
    assert minilm.get("/difficulty/info").json()["semantic_scorer"] == "minilm:abcd1234"


def test_info_shows_unavailable(real_app, monkeypatch):
    monkeypatch.setenv("EDUSYNC_C2_SEMANTIC", "minilm")
    monkeypatch.setattr(RT, "get_default_encoder", lambda model_dir: None)
    r = TestClient(real_app).get("/difficulty/info")
    assert r.json()["semantic_scorer"] == "minilm:unavailable"


def test_pipeline_info_field():
    assert DifficultyPipeline().info().semantic_scorer == "off"
    pipe = DifficultyPipeline(semantic_status="minilm:cafe1234")
    assert pipe.info().semantic_scorer == "minilm:cafe1234"


def test_off_has_no_semantic_warning(real_app):
    body = post(TestClient(real_app)).json()
    assert not [w for w in body["warnings"] if "semantic" in w]
    assert feature(body, "expl.semantic_support_score")["source"] == "semantic:off"


def test_unavailable_warns_once_and_keeps_status_200(real_app, monkeypatch):
    monkeypatch.setenv("EDUSYNC_C2_SEMANTIC", "minilm")
    monkeypatch.setattr(RT, "get_default_encoder", lambda model_dir: None)
    r = post(TestClient(real_app))
    assert r.status_code == 200
    body = r.json()
    assert body["warnings"].count(SEMANTIC_UNAVAILABLE_WARNING) == 1
    assert SEMANTIC_UNAVAILABLE_WARNING == (
        "semantic scorer unavailable: semantic features are unavailable"
    )
    for name in SEMANTIC:
        f = feature(body, name)
        assert (f["status"], f["value"], f["source"]) == ("unavailable", None, "semantic:minilm")
        assert f["note"] == "semantic scorer unavailable"


def test_available_scorer_has_no_warning(minilm):
    body = post(minilm).json()
    assert not [w for w in body["warnings"] if "semantic" in w]


# ------------------------------------------------------------------ end to end


def test_end_to_end_counts_and_determinism(minilm):
    first = post(minilm)
    assert first.status_code == 200
    body = first.json()
    assert feature(body, "expl.semantic_definition_count")["value"] == 1.0
    assert feature(body, "expl.semantic_example_count")["value"] == 1.0
    assert feature(body, "expl.semantic_support_score")["value"] == 1.0
    assert feature(body, "expl.semantic_support_score")["source"] == "semantic:minilm"
    assert post(minilm).json() == body


def test_baseline_output_is_unchanged_by_the_scorer(real_app, monkeypatch, keyword_encoder):
    off = post(TestClient(real_app)).json()["predictions"]
    get_pipeline.cache_clear()
    monkeypatch.setenv("EDUSYNC_C2_SEMANTIC", "minilm")
    monkeypatch.setattr(RT, "get_default_encoder", lambda model_dir: keyword_encoder)
    on = post(TestClient(real_app)).json()["predictions"]
    assert [(p["score"], p["difficulty"]) for p in off] == [
        (p["score"], p["difficulty"]) for p in on
    ]


def test_logs_never_contain_slide_text(minilm, caplog):
    caplog.set_level(logging.DEBUG, logger="edusync_c2")
    assert post(minilm).status_code == 200
    logged = " ".join(r.getMessage() for r in caplog.records)
    assert logged
    for word in SECRET.split():
        assert word not in logged
    assert "DEFN" not in logged
    assert "EXMP" not in logged


# ------------------------------------------------------------------ failures


@pytest.mark.parametrize("mode", ["raise", "short"])
def test_scoring_failure_is_a_503_without_leaks(
    real_app, monkeypatch, keyword_encoder, caplog, mode
):
    monkeypatch.setenv("EDUSYNC_C2_SEMANTIC", "minilm")
    broken = PrototypesOnly(keyword_encoder, mode)
    monkeypatch.setattr(RT, "get_default_encoder", lambda model_dir: broken)
    caplog.set_level(logging.DEBUG, logger="edusync_c2")
    r = post(TestClient(real_app))
    assert r.status_code == 503
    assert r.json() == {"detail": "semantic scorer failed"}
    logged = " ".join(m.getMessage() for m in caplog.records)
    assert "rejected status=503 reason=semantic_scorer" in logged
    for word in SECRET.split():
        assert word not in logged
        assert word not in r.text


def test_the_next_request_is_served_again_after_a_failure(real_app, monkeypatch, keyword_encoder):
    monkeypatch.setenv("EDUSYNC_C2_SEMANTIC", "minilm")
    broken = PrototypesOnly(keyword_encoder, "raise")
    current = {"encoder": broken}
    monkeypatch.setattr(RT, "get_default_encoder", lambda model_dir: current["encoder"])
    client = TestClient(real_app)
    assert post(client).status_code == 503
    broken.mode = "ok"  # the same scorer object is reused; the encoder recovers
    assert post(client).status_code == 200


def test_prototype_encoding_failure_makes_the_scorer_unavailable(
    real_app, monkeypatch, keyword_encoder, caplog
):
    class NoPrototypes:
        dims = 4
        sha = "feedbeef"

        def encode(self, texts):
            raise RuntimeError("no weights")

    monkeypatch.setenv("EDUSYNC_C2_SEMANTIC", "minilm")
    monkeypatch.setattr(RT, "get_default_encoder", lambda model_dir: NoPrototypes())
    caplog.set_level(logging.INFO, logger="edusync_c2")
    client = TestClient(real_app)
    assert client.get("/difficulty/info").json()["semantic_scorer"] == "minilm:unavailable"
    r = post(client)
    assert r.status_code == 200
    assert SEMANTIC_UNAVAILABLE_WARNING in r.json()["warnings"]
    assert "semantic_unavailable reason=prototype_encoding" in caplog.text


def test_scorer_wraps_encoder_errors_without_their_text(keyword_encoder):
    scorer = SemanticScorer(keyword_encoder)
    scorer._shared.encoder = PrototypesOnly(keyword_encoder, "raise")
    with pytest.raises(SemanticScoringError) as info:
        scorer.score_sentences([f"a sentence about {SECRET} here"])
    assert str(info.value) == "semantic scorer failed"
    assert SECRET not in str(info.value)
    assert isinstance(info.value.__cause__, RuntimeError)
    assert info.value.status_code == 503
