"""The heuristic concept filter wired through the pipeline, plus the default-path parity checks."""

import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

import edusync_c2.difficulty.pipeline as pipeline_mod
from edusync_c2.difficulty import DifficultyPipeline
from edusync_c2.difficulty.concept_filter import HeuristicConceptFilter
from edusync_c2.difficulty.features.language import language_features
from edusync_c2.difficulty.fusion import build_fusion_input
from edusync_c2.difficulty.graph import GraphView
from edusync_c2.difficulty.schemas import AlignmentStatus, C1GraphResponse, ConceptRecord
from edusync_c2.difficulty.synthetic import build_c1_graph, build_pptx, make_deck_spec

SRC_DIR = Path(__file__).resolve().parents[2] / "src"
FILENAME = "tde_lecture.pptx"
SLIDES = [
    ("Best Practices", ["Best Practices for storing records in a database."]),
    (
        "Transparent Data Encryption (TDE)",
        [
            "TDE protects data at rest on the disk.",
            "Transparent Data Encryption uses a master key stored apart.",
        ],
    ),
    ("Structured Query Language (SQL)", ["SQL is used to query tables.", "Learn SQL by practice."]),
]


def _node(cid, label, freq=1):
    return {"id": cid, "label": label, "frequency": freq}


def _edge(src, tgt, conf=0.6, kind="inferred"):
    return {
        "source": src,
        "target": tgt,
        "relation_type": "prerequisite",
        "direction": "source_is_prerequisite_of_target",
        "confidence": conf,
        "source_type": kind,
    }


def _graph() -> C1GraphResponse:
    nodes = [
        _node("tde", "TDE", 4),
        _node("transparent-data-encryption", "Transparent Data Encryption", 2),
        _node("sql", "SQL", 5),
        _node("structured-query-language", "Structured Query Language", 1),
        _node("sql-lower", "sql", 1),
        _node("best-practices", "Best Practices", 1),
    ]
    edges = [
        _edge("sql", "tde", 0.5),  # into an alias: rewritten
        _edge("structured-query-language", "transparent-data-encryption", 0.9),  # duplicate
        _edge("sql", "sql-lower"),  # self-loop after the merge
        _edge("best-practices", "sql"),  # touches a dropped concept: cut
    ]
    return C1GraphResponse.model_validate(
        {"lecture_id": "lec-tde", "graph": {"nodes": nodes, "edges": edges}, "concepts": []}
    )


@pytest.fixture(scope="module")
def deck(tmp_path_factory) -> bytes:
    prs = Presentation()
    for title, lines in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = "\n".join(lines)
    path = tmp_path_factory.mktemp("deck") / "d.pptx"
    prs.save(str(path))
    return path.read_bytes()


def heuristic():
    return DifficultyPipeline(concept_filter=HeuristicConceptFilter())


def stages(resp):
    return {s.name: s for s in resp.stages}


def test_heuristic_end_to_end(deck):
    resp = heuristic().run(deck, FILENAME, _graph())
    by_id = {p.concept_id: p for p in resp.predictions}
    assert sorted(by_id) == ["structured-query-language", "transparent-data-encryption"]
    assert by_id["transparent-data-encryption"].merged_from == ["tde"]
    assert by_id["structured-query-language"].merged_from == ["sql", "sql-lower"]
    assert "best-practices" not in by_id
    st = stages(resp)
    assert st["concepts"].status == "ok"
    assert "filter heuristic: 3 merged, 1 dropped" in st["concepts"].detail
    assert "3 aliases merged, 1 edges rewritten, 1 edges dropped" in st["graph"].detail
    assert any("became self-loops" in w for w in resp.warnings)
    assert len(resp.warnings) == len(set(resp.warnings))
    tde = by_id["transparent-data-encryption"]
    direct = next(f for f in tde.features if f.name == "prereq.direct_count")
    assert direct.value == 1.0  # the duplicate edges collapsed into one


def test_determinism_and_log_level_independence(deck, caplog):
    caplog.set_level("DEBUG", logger="edusync_c2")
    first = heuristic().run(deck, FILENAME, _graph()).model_dump_json()
    caplog.set_level("WARNING", logger="edusync_c2")
    second = heuristic().run(deck, FILENAME, _graph()).model_dump_json()
    assert first == second


def test_log_order_and_merged_suffix(deck, caplog):
    caplog.set_level("INFO", logger="edusync_c2")
    heuristic().run(deck, FILENAME, _graph())
    msgs = [r.getMessage() for r in caplog.records if r.name == "edusync_c2.difficulty.pipeline"]
    assert [m.split()[0] for m in msgs] == [
        "materials",
        "graph",
        "concepts",
        "features",
        "fusion",
        "model",
        "done",
    ]
    assert re.fullmatch(
        r"graph source=real nodes=6 edges=4 prerequisite_edges=4 ignored=0", msgs[1]
    )
    assert re.fullmatch(
        r"concepts total=2 aligned=2 not_found=0 filter=heuristic dropped=1 merged=3", msgs[2]
    )
    assert msgs[4].startswith("fusion inputs=2 ")


def test_privacy_no_slide_text_in_logs(deck, caplog):
    caplog.set_level("DEBUG", logger="edusync_c2")
    heuristic().run(deck, FILENAME, _graph())
    texts = [ln for _, lines in SLIDES for ln in lines if len(ln) > 25]
    assert texts
    assert any(r.levelname == "DEBUG" for r in caplog.records)
    for r in caplog.records:
        if r.name.startswith("edusync_c2"):
            m = r.getMessage()
            assert FILENAME not in m
            assert not any(t in m for t in texts)


def _spy_graph_view(monkeypatch):
    built = []
    real = GraphView

    def spy(*a, **k):
        built.append(k)
        return real(*a, **k)

    monkeypatch.setattr(pipeline_mod, "GraphView", spy)
    return built


def test_default_path_builds_one_plain_view(monkeypatch):
    spec = make_deck_spec(7, 4, "lec-parity")
    graph = C1GraphResponse.model_validate(build_c1_graph(spec, "v1"))
    built = _spy_graph_view(monkeypatch)
    resp = DifficultyPipeline().run(build_pptx(spec), "d.pptx", graph)
    assert built == [{}]  # the raw view is reused: no second view with aliases or drops
    assert all(p.merged_from == [] for p in resp.predictions)


def test_heuristic_without_changes_keeps_the_raw_view(monkeypatch):
    spec = make_deck_spec(7, 4, "lec-parity")
    graph = C1GraphResponse.model_validate(build_c1_graph(spec, "v1"))
    built = _spy_graph_view(monkeypatch)
    resp = heuristic().run(build_pptx(spec), "d.pptx", graph)
    assert built == [{}]
    assert "filter heuristic" not in stages(resp)["concepts"].detail
    assert "merged" not in stages(resp)["graph"].detail


def test_hash_seed_independence(deck, tmp_path):
    path = tmp_path / "d.pptx"
    path.write_bytes(deck)
    gpath = tmp_path / "g.json"
    gpath.write_text(_graph().model_dump_json(), encoding="utf-8")
    code = (
        "import sys, hashlib\n"
        "from edusync_c2.difficulty import DifficultyPipeline\n"
        "from edusync_c2.difficulty.concept_filter import HeuristicConceptFilter\n"
        "from edusync_c2.difficulty.schemas import C1GraphResponse\n"
        "g = C1GraphResponse.model_validate_json(open(sys.argv[2]).read())\n"
        "r = DifficultyPipeline(concept_filter=HeuristicConceptFilter())"
        ".run(open(sys.argv[1], 'rb').read(), 'd.pptx', g)\n"
        "print(hashlib.sha256(r.model_dump_json().encode()).hexdigest())\n"
    )
    out = []
    for seed in ("0", "1"):
        done = subprocess.run(
            [sys.executable, "-c", code, str(path), str(gpath)],
            capture_output=True,
            text=True,
            env={
                **os.environ,
                "PYTHONHASHSEED": seed,
                "PYTHONPATH": os.pathsep.join([str(SRC_DIR), *sys.path]),
            },
            check=True,
        )
        out.append(done.stdout.strip())
    assert out[0] == out[1]


# ---------------------------------------------------------------- features with alias labels


def _concept(label, aliases=()):
    return ConceptRecord(
        id="c",
        label=label,
        alignment_status=AlignmentStatus.ALIGNED,
        alias_labels=list(aliases),
        merged_from=["x"] if aliases else [],
    )


def _density(make_slide, concept, text):
    feats = language_features(concept, [make_slide(1, text)])
    return next(f.value for f in feats if f.name == "lang.concept_term_density")


def _n_words(text):
    return len(re.findall(r"\w+", text))


def test_density_counts_alias_mentions_once(make_slide):
    text = "TDE helps. Transparent  Data Encryption works. tde again and the word tdes."
    c = _concept("Transparent Data Encryption", ["TDE"])
    assert _density(make_slide, c, text) == pytest.approx(100.0 * 3 / _n_words(text))


def test_density_longest_first_and_distinct(make_slide):
    text = "Data Encryption is part of Transparent Data Encryption."
    c = _concept("Transparent Data Encryption", ["Data Encryption", "data  encryption", "  "])
    assert _density(make_slide, c, text) == pytest.approx(100.0 * 2 / _n_words(text))


def test_density_without_aliases_is_the_old_path(make_slide):
    text = "Transparent  Data Encryption appears; Transparent Data Encryption twice."
    c = _concept("Transparent Data Encryption")
    assert _density(make_slide, c, text) == pytest.approx(100.0 * 1 / _n_words(text))


def test_density_blank_terms_only(make_slide):
    c = _concept(" ", ["  "])
    assert _density(make_slide, c, "some text here") == 0.0


def test_fusion_text_mentions_aliases(make_slide):
    s = make_slide(1, "body text")
    plain = build_fusion_input(_concept("TDE"), [], [s])
    merged = build_fusion_input(_concept("Transparent Data Encryption", ["TDE", "T.D.E."]), [], [s])
    assert plain.text == "TDE. \nbody text"
    assert merged.text.startswith("Transparent Data Encryption.  Also known as: TDE, T.D.E..\n")


def test_drop_without_edges_adds_no_graph_note(deck):
    graph = C1GraphResponse.model_validate(
        {
            "graph": {
                "nodes": [_node("best-practices", "Best Practices"), _node("sql", "SQL")],
                "edges": [],
            }
        }
    )
    resp = heuristic().run(deck, FILENAME, graph)
    assert [p.concept_id for p in resp.predictions] == ["sql"]
    assert "merged" not in stages(resp)["graph"].detail
    assert "filter heuristic: 0 merged, 1 dropped" in stages(resp)["concepts"].detail
