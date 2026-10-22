"""Deterministic synthetic lecture data for testing the difficulty pipeline.

SYNTHETIC DATA: no real course material is used. The difficulty labels come from a transparent
rule over the same properties the generated decks expose (``synthetic_label``), so any model
score measured against them is circular and must never be reported as accuracy.

Everything is derived from a seed: same seed, same spec, same C1 graphs, same labels. (PPTX files
embed zip timestamps, so compare parsed content, not bytes.) All text is ASCII so the PDF writer
needs no font encoding.
"""

import csv
import random
import struct
import zlib
from dataclasses import dataclass
from io import BytesIO, StringIO
from typing import Literal

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt
from pydantic import BaseModel, Field

from edusync_c2.difficulty.ids import slug_id
from edusync_c2.difficulty.schemas import DifficultyLabel


def slugify(text: str) -> str:
    """C1 v1 id rule: lowercase letters/digits joined by single hyphens (ASCII labels)."""
    return slug_id(text)


_ADJECTIVES = [
    "Linear", "Cyclic", "Sparse", "Dense", "Static", "Dynamic",
    "Modular", "Nested", "Stable", "Layered", "Bounded", "Ordered",
]  # fmt: skip
_NOUNS = [
    "Signal", "Filter", "Buffer", "Kernel", "Matrix", "Queue",
    "Tensor", "Router", "Sampler", "Encoder", "Scheduler", "Index",
]  # fmt: skip
_FILLER = [
    "the", "model", "keeps", "each", "stage", "simple", "while", "every", "part", "stays",
    "stable", "and", "small", "steps", "remain", "clear", "across", "whole", "order", "work",
]  # fmt: skip
_FORMULAS = [
    "L(w) = sum((y - w*x)^2) / n",
    "y = a*x^2 + b*x + c",
    "E = m * (c^2)",
    "z = (x - mu) / sigma",
    "\\frac{a}{b} = (a + b) / 2",
]
_CODE = ["def run(item):", "    result = item.value;", "    return result;"]
_MAX_FILLER_LINE = 10  # words per filler paragraph


# ------------------------------------------------------------------------------------- specs


class ConceptSpec(BaseModel):
    """Properties of one synthetic concept slide; the label rule reads exactly these."""

    label: str
    definition: bool
    examples: int = Field(ge=0)
    steps: int = Field(ge=0)
    words: int = Field(ge=0)  # target number of body words
    formula: bool
    code: bool
    picture: bool
    bullet_depth: int = Field(ge=0)  # deepest bullet level on the slide
    prerequisites: list[str] = Field(default_factory=list)  # labels of earlier concepts

    @property
    def concept_id(self) -> str:
        """C1 v0 id convention: the casefolded, stripped label."""
        return self.label.strip().casefold()

    @property
    def slug(self) -> str:
        """C1 v1 id convention: lowercase letters and digits joined by single hyphens."""
        return slugify(self.label)


class DeckSpec(BaseModel):
    lecture_id: str
    title: str
    concepts: list[ConceptSpec]


def make_deck_spec(seed: int, n_concepts: int = 8, lecture_id: str | None = None) -> DeckSpec:
    """Build a random but reproducible deck spec; prerequisites only point to earlier concepts."""
    labels_pool = [f"{a} {n}" for a in _ADJECTIVES for n in _NOUNS]
    if not 1 <= n_concepts <= len(labels_pool):
        raise ValueError(f"n_concepts must be between 1 and {len(labels_pool)}")
    rng = random.Random(seed)
    labels = rng.sample(labels_pool, n_concepts)
    concepts = []
    for i, label in enumerate(labels):
        n_prereq = min(i, rng.choice([0, 1, 1, 2, 3]))
        concepts.append(
            ConceptSpec(
                label=label,
                definition=rng.random() < 0.6,
                examples=rng.randint(0, 3),
                steps=rng.choice([0, 0, 2, 3, 4]),
                words=rng.randint(10, 120),
                formula=rng.random() < 0.3,
                code=rng.random() < 0.2,
                picture=rng.random() < 0.4,
                bullet_depth=rng.randint(0, 3),
                prerequisites=rng.sample(labels[:i], n_prereq),
            )
        )
    return DeckSpec(
        lecture_id=lecture_id or f"synthetic-{seed}",
        title=f"Synthetic Lecture {seed}",
        concepts=concepts,
    )


# ----------------------------------------------------------------------------- slide content


def _slide_lines(deck: DeckSpec, c: ConceptSpec) -> list[tuple[str, int]]:
    """Body lines (text, bullet level) of one concept slide, built from the spec."""
    lines: list[tuple[str, int]] = []
    if c.definition:
        lines.append((f"{c.label} is defined as a synthetic construct for pipeline tests.", 0))
    lines += [(f"For example, {c.label} appears in case {k}.", 0) for k in range(1, c.examples + 1)]
    lines += [(f"Step {k}: apply {c.label} to the next item.", 0) for k in range(1, c.steps + 1)]
    lines += [(f"Detail at level {lv}.", lv) for lv in range(1, c.bullet_depth + 1)]
    rng = random.Random(f"{deck.lecture_id}:{c.label}")
    missing = c.words - sum(len(text.split()) for text, _ in lines)
    while missing > 0:
        take = min(missing, _MAX_FILLER_LINE)
        lines.append((" ".join(rng.choice(_FILLER) for _ in range(take)) + ".", 0))
        missing -= take
    if c.formula:
        lines.append((rng.choice(_FORMULAS), 0))
    if c.code:
        lines += [(text, 0) for text in _CODE]
    return lines


def _slide_texts(deck: DeckSpec) -> list[tuple[str, list[tuple[str, int]]]]:
    """(title, body lines) for every slide: title slide, one per concept, closing slide."""
    slides = [(deck.title, [("Synthetic lecture data for pipeline testing.", 0)])]
    slides += [(c.label, _slide_lines(deck, c)) for c in deck.concepts]
    slides.append(("Questions?", []))
    return slides


def _png(size: int = 8, rgb: tuple[int, int, int] = (70, 110, 170)) -> bytes:
    """A tiny solid-colour PNG, generated without any imaging library."""

    def chunk(kind: bytes, payload: bytes) -> bytes:
        body = kind + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))

    row = b"\x00" + bytes(rgb) * size
    header = struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(row * size))
        + chunk(b"IEND", b"")
    )


# ------------------------------------------------------------------------------------- files


_FONT_SIZES_PT = (10, 12, 14, 18, 24)
_Box = tuple[float, float, float, float]  # left, top, width, height in inches


@dataclass(frozen=True)
class _SlideLayout:
    """Geometry of one concept slide (inches, points); the default is the classic fixed layout."""

    text_box: _Box = (0.5, 1.5, 6.0, 5.0)
    font_pt: int = 14
    picture_width: float = 2.0
    shapes: tuple[_Box, ...] = ()  # autoshapes
    connectors: tuple[_Box, ...] = ()  # x1, y1, dx, dy in inches
    table: tuple[int, int, _Box] | None = None  # rows, cols, frame


def _draw_box(rng: random.Random) -> _Box:
    return (
        rng.uniform(0.3, 8.0),
        rng.uniform(1.2, 6.5),
        rng.uniform(0.4, 2.5),
        rng.uniform(0.3, 1.5),
    )


def _draw_layout(rng: random.Random) -> _SlideLayout:
    """Draw one slide's geometry from the layout stream, always in the same order."""
    width = rng.uniform(3.0, 9.0)
    height = rng.uniform(1.5, 5.5)
    if rng.random() < 0.05:  # rare overflow: the box runs off the right edge of the slide
        width += 3.0
    font_pt = rng.choice(_FONT_SIZES_PT)
    picture_width = rng.uniform(1.0, 3.5)
    shapes = tuple(_draw_box(rng) for _ in range(rng.randint(0, 6)))
    connectors = tuple(_draw_box(rng) for _ in range(rng.randint(0, 3)))
    has_table = rng.random() < 0.25
    rows, cols = rng.randint(2, 4), rng.randint(2, 4)
    frame = (0.5, rng.uniform(4.5, 6.0), rng.uniform(2.0, 6.0), rng.uniform(0.8, 1.5))
    return _SlideLayout(
        text_box=(0.5, 1.5, width, height),
        font_pt=font_pt,
        picture_width=picture_width,
        shapes=shapes,
        connectors=connectors,
        table=(rows, cols, frame) if has_table else None,
    )


def _add_extras(slide, layout: _SlideLayout) -> None:
    """Diagram elements and the optional table of a varied slide."""
    for left, top, width, height in layout.shapes:
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
    for x1, y1, dx, dy in layout.connectors:
        slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x1 + dx), Inches(y1 + dy)
        )
    if layout.table is not None:
        rows, cols, (left, top, width, height) = layout.table
        slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))


def build_pptx(spec: DeckSpec, *, vary_layout: bool = False) -> bytes:
    """One slide per concept plus a title slide and a final 'Questions?' slide.

    ``vary_layout`` (opt-in) varies box sizes, font sizes, picture size, diagram elements and
    tables per concept slide from a separate ``{lecture_id}:geom`` stream, so text, labels and
    graphs are unaffected. The default reproduces the fixed layout exactly.
    """
    prs = Presentation()
    concepts = {c.label: c for c in spec.concepts}
    rng = random.Random(f"{spec.lecture_id}:geom") if vary_layout else None
    for position, (title, lines) in enumerate(_slide_texts(spec)):
        if position == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title
            slide.placeholders[1].text = lines[0][0]
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        layout = _draw_layout(rng) if rng is not None and title in concepts else _SlideLayout()
        if lines:
            left, top, width, height = layout.text_box
            frame = slide.shapes.add_textbox(
                Inches(left), Inches(top), Inches(width), Inches(height)
            ).text_frame
            frame.word_wrap = True
            for n, (text, level) in enumerate(lines):
                paragraph = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.level = level
                paragraph.font.size = Pt(layout.font_pt)
        if title in concepts and concepts[title].picture:
            slide.shapes.add_picture(
                BytesIO(_png()), Inches(7), Inches(1.5), Inches(layout.picture_width)
            )
        _add_extras(slide, layout)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_text_pdf(pages: list[list[str]]) -> bytes:
    """Hand-written text-only PDF, one page per list of lines (an empty list is a blank page)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")  # object 3
    kids = []
    for lines in pages:
        page_no = len(objects) + 1
        kids.append(f"{page_no} 0 R")
        shown = " T* ".join(f"({_pdf_escape(line)}) Tj" for line in lines)
        stream = f"BT /F1 12 Tf 72 720 Td 14 TL {shown} ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {page_no + 1} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1", errors="replace")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def build_pdf(spec: DeckSpec) -> bytes:
    """Text-only PDF with one page per slide (title first, then the body lines)."""
    pages = [[title, *(text for text, _ in lines)] for title, lines in _slide_texts(spec)]
    return build_text_pdf(pages)


# -------------------------------------------------------------------------------------- graph


def _pairs(spec: DeckSpec, use_slug: bool = False) -> list[tuple[str, str, int]]:
    """Co-mentioned (source_id, target_id, page) triples: neighbours and prerequisite pairs."""

    def ident(c: ConceptSpec) -> str:
        return c.slug if use_slug else c.concept_id

    ids = {c.label: ident(c) for c in spec.concepts}
    found: dict[tuple[str, str], int] = {}
    for i, c in enumerate(spec.concepts):
        page = i + 2  # slide 1 is the title slide
        related = [ids[p] for p in c.prerequisites]
        if i:
            related.append(ident(spec.concepts[i - 1]))
        for other in related:
            found.setdefault(tuple(sorted((ident(c), other))), page)
    return [(a, b, page) for (a, b), page in sorted(found.items())]


def build_c1_graph(spec: DeckSpec, version: Literal["v0", "v1"] = "v1") -> dict:
    """C1 response shape. v0: co-mention edges only; v1 adds directed prerequisite edges.

    Ids follow the C1 convention of each version: v0 casefolded labels, v1 hyphenated slugs.
    """
    use_slug = version == "v1"

    def ident(c: ConceptSpec) -> str:
        return c.slug if use_slug else c.concept_id

    labels = {ident(c): c.label for c in spec.concepts}
    nodes = []
    for c in spec.concepts:
        node = {"id": ident(c), "label": c.label}
        if version == "v1":
            node["description"] = f"{c.label}: synthetic concept used for pipeline testing."
        node["frequency"] = 1 + c.examples + c.steps
        nodes.append(node)
    edges: list[dict] = [
        {
            "source": a,
            "target": b,
            "relation_type": "co_mentioned",
            "evidence_count": 1,
            "evidence": [
                {"page": page, "sentence": f"{labels[a]} and {labels[b]} appear together."}
            ],
        }
        for a, b, page in _pairs(spec, use_slug)
    ]
    if version == "v1":
        for i, c in enumerate(spec.concepts):
            for prerequisite in c.prerequisites:
                source = slugify(prerequisite)
                rng = random.Random(f"{spec.lecture_id}:{source}->{ident(c)}")
                edges.append(
                    {
                        "source": source,
                        "target": ident(c),
                        "relation_type": "prerequisite",
                        "direction": "source_is_prerequisite_of_target",
                        "confidence": round(0.5 + 0.5 * rng.random(), 2),
                        "source_type": "curriculum_defined" if rng.random() < 0.4 else "inferred",
                        "evidence_count": 1,
                        "evidence": [
                            {"page": i + 2, "sentence": f"{prerequisite} comes before {c.label}."}
                        ],
                    }
                )
    note = "Synthetic co-mention graph." if version == "v0" else "Synthetic prerequisite graph."
    return {
        "lecture_id": spec.lecture_id,
        "filename": f"{spec.lecture_id}.pptx",
        "page_count": len(spec.concepts) + 2,
        "concept_count": len(spec.concepts),
        "concepts": [
            {"name": c.label, "frequency": 1 + c.examples + c.steps} for c in spec.concepts
        ],
        "graph": {
            "node_count": len(nodes),
            "edge_count": len(edges),
            "nodes": nodes,
            "edges": edges,
            "note": note,
        },
        "message": "Candidate knowledge graph generated successfully.",
    }


# ------------------------------------------------------------------------------------ labels

LOW_BELOW = 1.0
MEDIUM_BELOW = 2.5


def synthetic_label(c: ConceptSpec, n_prerequisites: int) -> DifficultyLabel:
    """SYNTHETIC rule label for pipeline testing only (circular, not a measure of difficulty).

    score = 0.8 * prerequisites + words / 50 + 1.5 * formula + 1.5 * code + 0.3 * bullet_depth
            - 0.8 * definition - 0.5 * min(examples, 3) - 0.25 * min(steps, 4) - 0.6 * picture

    Low if score < 1.0, Medium if score < 2.5, otherwise High.
    """
    score = (
        0.8 * n_prerequisites
        + c.words / 50
        + 1.5 * c.formula
        + 1.5 * c.code
        + 0.3 * c.bullet_depth
        - 0.8 * c.definition
        - 0.5 * min(c.examples, 3)
        - 0.25 * min(c.steps, 4)
        - 0.6 * c.picture
    )
    if score < LOW_BELOW:
        return DifficultyLabel.LOW
    return DifficultyLabel.MEDIUM if score < MEDIUM_BELOW else DifficultyLabel.HIGH


def build_labels_csv(specs: list[DeckSpec]) -> str:
    """CSV of rule labels; the ``synthetic`` column is always ``true``."""
    buf = StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(
        [
            "lecture_id",
            "concept_id",
            "concept_slug",
            "concept_label",
            "difficulty_label",
            "synthetic",
        ]
    )
    for spec in specs:
        for c in spec.concepts:
            label = synthetic_label(c, len(c.prerequisites))
            row = [spec.lecture_id, c.concept_id, c.slug, c.label, label.value, "true"]
            writer.writerow(row)
    return buf.getvalue()
