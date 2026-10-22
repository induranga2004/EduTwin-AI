"""Lecture materials parser for the difficulty pipeline (PPTX and PDF).

C2 owns its parser because C1's extractor flattens text and loses titles, speaker notes, bullet
levels, pictures, charts and tables. A PPTX keeps all of that; a PDF only yields text per page,
so its structure fields stay ``None`` and ``structure_available`` is ``False``.
"""

import re
from io import BytesIO

from lxml.etree import LxmlError
from pptx import Presentation
from pptx.exc import PythonPptxError
from pptx.shapes.autoshape import Shape
from pptx.shapes.connector import Connector
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.shapes.placeholder import PlaceholderPicture
from pptx.slide import Slide
from pptx.text.text import Font, TextFrame
from pypdf import PdfReader

from edusync_c2.difficulty.geometry import (
    Box,
    GroupFrame,
    Rect,
    apply_rotation,
    to_box,
    to_slide_space,
    union_area,
)
from edusync_c2.difficulty.schemas import MaterialsDoc, ParagraphRecord, SlideRecord

ALLOWED_EXTENSIONS = (".pptx", ".pdf")
PDF_STRUCTURE_WARNING = "PDF: structure fields (pictures, charts, bullet levels) are unavailable"


class MaterialsError(Exception):
    """A rejected upload, carrying the HTTP status the API should answer with."""

    def __init__(self, status_code: int, message: str) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.message = message


# ------------------------------------------------------------------------------------ detectors

_LATEX = re.compile(r"\\(?:frac|sum|int|sqrt)(?![A-Za-z])|\$\$")
_MATH_SYMBOLS = ("\u2211", "\u222b", "\u221a", "\u2264", "\u2265", "\u2202", "\u221e", "\u00b1")
_COMPARISONS = re.compile(r"==|=>|<=|>=|!=")
_EQUATION = re.compile(r"[\w)\]]\s*=\s*[\w(\-]")
_LONG_WORD = re.compile(r"[A-Za-z]{4,}")
_KEYWORDS = re.compile(r"\b(?:def|class|import|function)\s+\w|\breturn[ ;]")
_SQL = re.compile(r"\bSELECT\b.*\bFROM\b", re.IGNORECASE)
_CODE_OPERATORS = ("==", "=>", "->")
_INDENTED = re.compile(r"^(?: {4}|\t)\S", re.MULTILINE)


def detect_formula(text: str) -> bool:
    """Heuristic: does the text contain a formula or equation?

    True when any of these holds:
    - a LaTeX marker (``\\frac``, ``\\sum``, ``\\int``, ``\\sqrt``, ``$$``);
    - a math symbol (Unicode sum, integral, root, less/greater-or-equal, partial, infinity,
      plus-minus) or the ASCII ``<=`` / ``>=``;
    - an equation-like line: a single ``=`` (not part of ``==``, ``=>``, ``<=``, ``>=``, ``!=``)
      between operands, with an operator, parenthesis or exponent on the line, at most one word
      of four or more letters (so prose sentences do not count), and not ending in ``;`` (code).
    """
    if _LATEX.search(text) or any(s in text for s in (*_MATH_SYMBOLS, "<=", ">=")):
        return True
    for line in text.splitlines():
        line = line.strip()
        if not line or line.endswith(";") or len(line) > 80:
            continue
        bare = _COMPARISONS.sub(" ", line)
        has_operator = any(c in bare for c in "+*/^()") or " - " in bare
        if _EQUATION.search(bare) and has_operator and len(_LONG_WORD.findall(bare)) <= 1:
            return True
    return False


def detect_code(text: str) -> bool:
    """Heuristic: does the text look like source code?

    True when at least two of these signals are present: a line ending in ``;``; a keyword
    (``def x``, ``class x``, ``import x``, ``function x``, ``return``, ``SELECT ... FROM``);
    braces; one of ``==``, ``=>``, ``->``; a line indented by four spaces or a tab.
    """
    signals = (
        any(line.rstrip().endswith(";") for line in text.splitlines()),
        bool(_KEYWORDS.search(text) or _SQL.search(text)),
        "{" in text or "}" in text,
        any(op in text for op in _CODE_OPERATORS),
        bool(_INDENTED.search(text)),
    )
    return sum(signals) >= 2


# ------------------------------------------------------------------------------------- parsing


def parse_materials(data: bytes, filename: str, *, max_mb: int = 100) -> MaterialsDoc:
    """Validate and parse an uploaded PPTX or PDF into a ``MaterialsDoc``.

    Raises ``MaterialsError``: 400 (unsupported type, empty, wrong signature, unreadable,
    encrypted PDF), 413 (larger than ``max_mb`` MB), 422 (no text on any slide or page).
    """
    name = filename.lower()
    if not name.endswith(ALLOWED_EXTENSIONS):
        raise MaterialsError(400, "Unsupported file type; upload a .pptx or .pdf file")
    if not data:
        raise MaterialsError(400, "Uploaded file is empty")
    if len(data) > max_mb * 1024 * 1024:
        raise MaterialsError(413, f"File exceeds the {max_mb} MB limit")
    if name.endswith(".pdf"):
        if not data.startswith(b"%PDF-"):
            raise MaterialsError(400, "File content is not a PDF")
        doc = _parse_pdf(data, filename)
    else:
        if not data.startswith(b"PK"):
            raise MaterialsError(400, "File content is not a PPTX")
        doc = _parse_pptx(data, filename)
    if not any(slide.text.strip() for slide in doc.slides):
        raise MaterialsError(422, "No extractable text found in the file")
    return doc


def _index_list(indexes: list[int]) -> str:
    return ", ".join(str(i) for i in indexes)


def _parse_pdf(data: bytes, filename: str) -> MaterialsDoc:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise MaterialsError(400, "Encrypted PDFs are not supported")
        texts = [(page.extract_text() or "").strip() for page in reader.pages]
    except MaterialsError:
        raise
    except Exception as exc:
        raise MaterialsError(400, "Could not read the PDF") from exc
    slides = [
        SlideRecord(
            index=i,
            text=text,
            paragraphs=[
                ParagraphRecord(text=line.strip()) for line in text.splitlines() if line.strip()
            ],
            formula_hint=detect_formula(text),
            code_hint=detect_code(text),
        )
        for i, text in enumerate(texts, start=1)
    ]
    warnings = [PDF_STRUCTURE_WARNING]
    blank = [s.index for s in slides if not s.text]
    if blank:
        warnings.append(f"PDF pages without extractable text: {_index_list(blank)}")
    return MaterialsDoc(
        filename=filename,
        file_type="pdf",
        slide_count=len(slides),
        structure_available=False,
        slides=slides,
        warnings=warnings,
    )


# Everything a malformed shape can raise while its geometry or font is read; such a shape is
# counted in ``geometry_unknown_count`` instead of failing the whole upload.
_READ_ERRORS = (
    ValueError,
    AttributeError,
    TypeError,
    KeyError,
    NotImplementedError,
    PythonPptxError,
    LxmlError,
)
_SMARTART_URI_SUFFIX = "/drawingml/2006/diagram"


def _group_frame(shape: GroupShape) -> GroupFrame | None:
    """The group's child-to-parent mapping, or ``None`` when its transform cannot be read."""
    try:
        xfrm = shape._element.grpSpPr.xfrm
        off, ext, ch_off, ch_ext = xfrm.off, xfrm.ext, xfrm.chOff, xfrm.chExt
        return GroupFrame(
            int(off.x),
            int(off.y),
            int(ext.cx),
            int(ext.cy),
            int(ch_off.x),
            int(ch_off.y),
            int(ch_ext.cx),
            int(ch_ext.cy),
        )
    except _READ_ERRORS:
        return None


def _size_pt(properties) -> float | None:
    """Explicit font size of an ``a:rPr`` / ``a:defRPr`` element in points, else ``None``."""
    if properties is None:
        return None
    size = Font(properties).size
    return None if size is None else round(size.pt, 2)


def _frame_min_font(frame: TextFrame) -> float | None:
    """Smallest explicit size among the frame's text runs.

    A run's own size wins; a run without one takes the paragraph-level size. Nothing is read
    from the master, layout or list styles, so an inherited size is never guessed.
    """
    sizes: list[float] = []
    for paragraph in frame.paragraphs:
        paragraph_properties = paragraph._p.pPr
        paragraph_size = _size_pt(
            None if paragraph_properties is None else paragraph_properties.defRPr
        )
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = _size_pt(run._r.rPr)
            size = paragraph_size if size is None else size
            if size is not None:
                sizes.append(size)
    return min(sizes, default=None)


def _is_smartart(shape) -> bool:
    return isinstance(shape, GraphicFrame) and shape._element.graphicData_uri.endswith(
        _SMARTART_URI_SUFFIX
    )


def _is_diagram_leaf(shape) -> bool:
    """Autoshape (not a text box or placeholder), connector or freeform."""
    if isinstance(shape, Connector):
        return True
    return isinstance(shape, Shape) and not shape.is_placeholder and not shape._element.is_textbox


class _SlideCollector:
    """Accumulates paragraphs, shape counts and geometry while walking a slide's shape tree."""

    def __init__(self) -> None:
        self.paragraphs: list[ParagraphRecord] = []
        self.pictures = self.charts = self.tables = self.shapes = 0
        self.diagrams = self.smartart = self.table_cells = self.unknown = 0
        self.min_font: float | None = None
        self.text_boxes: list[Box] = []
        self.picture_boxes: list[Box] = []
        self.visual_boxes: list[Box] = []

    def add_frame(self, frame: TextFrame) -> None:
        for paragraph in frame.paragraphs:
            text = paragraph.text.replace("\v", " ").rstrip()
            if text.strip():
                self.paragraphs.append(ParagraphRecord(text=text, level=paragraph.level))

    def measure(self, shape, frames: list[GroupFrame | None]) -> None:
        """Record the shape's counts, font size and rect (title included, groups excluded).

        Any unreadable value counts the shape in ``unknown`` and records none of its geometry.
        """
        is_picture = isinstance(shape, Picture | PlaceholderPicture)
        is_smartart = _is_smartart(shape)
        is_diagram = is_smartart or _is_diagram_leaf(shape)
        self.smartart += is_smartart
        self.diagrams += is_diagram
        try:
            if shape.has_table:
                table = shape.table
                cells = len(table.rows) * len(table.columns)
                text_frames = [c.text_frame for row in table.rows for c in row.cells]
            elif shape.has_text_frame:
                cells = 0
                text_frames = [shape.text_frame]
            else:
                cells, text_frames = 0, []
            has_text = any(f.text.strip() for f in text_frames)
            fonts = [f for f in map(_frame_min_font, text_frames) if f is not None]
            rect = self._rect(shape, frames) if (is_picture or is_diagram or has_text) else None
        except _READ_ERRORS:
            self.unknown += 1
            return
        self.table_cells += cells
        if fonts:
            smallest = min(fonts)
            self.min_font = smallest if self.min_font is None else min(self.min_font, smallest)
        if rect is None:
            return
        box = to_box(rect)
        if has_text:
            self.text_boxes.append(box)
        if is_picture:
            self.picture_boxes.append(box)
        if is_picture or is_diagram:
            self.visual_boxes.append(box)

    @staticmethod
    def _rect(shape, frames: list[GroupFrame | None]) -> Rect:
        if any(frame is None for frame in frames):
            raise ValueError("enclosing group transform unreadable")
        x, y, w, h = shape.left, shape.top, shape.width, shape.height
        if x is None or y is None or w is None or h is None:
            raise ValueError("shape has no resolvable geometry")
        rect = apply_rotation((int(x), int(y), int(w), int(h)), shape.rotation)
        return to_slide_space(rect, frames)

    def walk(self, shapes, title_id: int | None, frames: list[GroupFrame | None] | None = None):
        frames = [] if frames is None else frames
        for shape in shapes:
            self.shapes += 1
            if shape.shape_id == title_id:
                self.measure(shape, frames)
                continue
            # isinstance, not shape.shape_type: that raises NotImplementedError for an autoshape
            # without recognisable geometry, which would turn a valid deck into an HTTP 400.
            if isinstance(shape, GroupShape):
                self.walk(shape.shapes, title_id, [*frames, _group_frame(shape)])
                continue
            self.measure(shape, frames)
            # A filled picture placeholder is a PlaceholderPicture (shape_type PLACEHOLDER).
            if isinstance(shape, Picture | PlaceholderPicture):
                self.pictures += 1
            if shape.has_chart:
                self.charts += 1
            if shape.has_table:
                self.tables += 1
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        self.paragraphs.append(ParagraphRecord(text=" | ".join(cells), level=0))
            elif shape.has_text_frame:
                self.add_frame(shape.text_frame)


def _parse_slide(slide: Slide, index: int, slide_w: int | None, slide_h: int | None) -> SlideRecord:
    collector = _SlideCollector()
    title_shape = slide.shapes.title
    title = None
    if title_shape is not None:
        collector.add_frame(title_shape.text_frame)
        title = title_shape.text_frame.text.strip() or None
    collector.walk(slide.shapes, title_shape.shape_id if title_shape is not None else None)
    text = "\n".join(p.text for p in collector.paragraphs)
    notes = ""
    if slide.has_notes_slide:
        notes = (slide.notes_slide.notes_text_frame.text or "").strip()
    areas: tuple[int | None, int | None, int | None] = (None, None, None)
    if slide_w is not None and slide_h is not None and slide_w > 0 and slide_h > 0:
        areas = (
            union_area(collector.text_boxes, slide_w, slide_h),
            union_area(collector.picture_boxes, slide_w, slide_h),
            union_area(collector.visual_boxes, slide_w, slide_h),
        )
    return SlideRecord(
        index=index,
        title=title,
        text=text,
        notes=notes,
        paragraphs=collector.paragraphs,
        picture_count=collector.pictures,
        chart_count=collector.charts,
        table_count=collector.tables,
        shape_count=collector.shapes,
        text_area_emu=areas[0],
        picture_area_emu=areas[1],
        visual_area_emu=areas[2],
        min_font_pt=collector.min_font,
        diagram_element_count=collector.diagrams,
        table_cell_count=collector.table_cells,
        smartart_count=collector.smartart,
        geometry_unknown_count=collector.unknown,
        formula_hint=detect_formula(text),
        code_hint=detect_code(text),
    )


def _parse_pptx(data: bytes, filename: str) -> MaterialsDoc:
    try:
        presentation = Presentation(BytesIO(data))
        width, height = presentation.slide_width, presentation.slide_height
        width = None if width is None else int(width)
        height = None if height is None else int(height)
        slides = [
            _parse_slide(s, i, width, height) for i, s in enumerate(presentation.slides, start=1)
        ]
    except Exception as exc:
        raise MaterialsError(400, "Could not read the PPTX") from exc
    warnings = []
    empty = [s.index for s in slides if not s.text.strip()]
    if empty:
        warnings.append(f"Slides without text: {_index_list(empty)}")
    return MaterialsDoc(
        filename=filename,
        file_type="pptx",
        slide_count=len(slides),
        structure_available=True,
        slide_width_emu=width,
        slide_height_emu=height,
        slides=slides,
        warnings=warnings,
    )
