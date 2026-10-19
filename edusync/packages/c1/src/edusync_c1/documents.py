"""PDF/PPTX text extraction: pure functions over bytes."""

import logging
from io import BytesIO

from pypdf import PdfReader

from edusync_c1.schemas import Page
from edusync_c1.validation import UploadError, validate_document

logger = logging.getLogger(__name__)


def extract_pdf_pages(content: bytes) -> list[Page]:
    """Extract text from each PDF page (encrypted PDFs raise a 400)."""
    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise UploadError(400, "Password-protected PDFs are not supported.")
    return [
        Page(page=number, text=(page.extract_text() or "").strip())
        for number, page in enumerate(reader.pages, start=1)
    ]


def extract_pptx_pages(content: bytes) -> list[Page]:
    """Extract text (and table rows) from each PowerPoint slide."""
    from pptx import Presentation

    try:
        presentation = Presentation(BytesIO(content))
    except Exception as exc:
        raise UploadError(400, "Could not read the PowerPoint presentation.") from exc

    pages: list[Page] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        pages.append(Page(page=number, text="\n".join(parts).strip()))
    return pages


def extract_document_pages(filename: str, content: bytes) -> list[Page]:
    """Validate then extract pages from a PDF or PPTX; failures map to 400/422."""
    extension = validate_document(filename, content)
    try:
        if extension == ".pdf":
            return extract_pdf_pages(content)
        return extract_pptx_pages(content)
    except UploadError:
        raise
    except Exception as exc:
        logger.exception("Document text extraction failed")
        raise UploadError(422, "Could not extract text from the lecture document.") from exc


def combine_page_text(pages: list[Page]) -> str:
    """Join non-empty pages as ``[Page N]`` blocks."""
    return "\n\n".join(f"[Page {p.page}]\n{p.text}" for p in pages if p.text).strip()
