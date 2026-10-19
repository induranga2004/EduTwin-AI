"""Shared fixtures: in-memory documents, a fake transcriber and an app/client."""

from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from edusync_c1.api import get_matcher, router
from edusync_c1.semantic import SemanticMatcher
from edusync_c1.transcription import get_transcriber
from edusync_core.settings import CoreSettings, get_settings

LECTURE_LINES = [
    "Gradient Descent minimizes a Loss Function.",
    "Gradient Descent uses a Learning Rate to update weights.",
    "The Loss Function measures prediction error.",
]


def make_text_pdf(lines: list[str]) -> bytes:
    """Build a one-page PDF with real text, with a valid xref table."""
    stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " T* ".join(f"({line}) Tj" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def make_blank_pdf(encrypt: bool = False) -> bytes:
    """A PDF with one blank page (no extractable text), optionally password-protected."""
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    if encrypt:
        writer.encrypt("secret")
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_pptx(slides: list[list[str]], table: list[list[str]] | None = None) -> bytes:
    """Build a PPTX with one text box per line on each slide, plus an optional table."""
    prs = Presentation()
    for lines in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for index, line in enumerate(lines):
            box = slide.shapes.add_textbox(Inches(1), Inches(1 + index), Inches(6), Inches(1))
            box.text_frame.text = line
    if table:
        shape = prs.slides[0].shapes.add_table(
            len(table), len(table[0]), Inches(1), Inches(5), Inches(6), Inches(1)
        )
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                shape.table.cell(r, c).text = value
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


class FakeTranscriber:
    """Returns canned segments; records calls."""

    def __init__(self, result: Any = None) -> None:
        self.calls: list[tuple[Path, str]] = []
        self.result = (
            result
            if result is not None
            else {
                "text": "",
                "language": "en",
                "segments": [
                    {"id": 0, "start": 0.0, "end": 4.0, "text": "Today we cover gradient descent."},
                    {"id": 1, "start": 4.0, "end": 9.0, "text": "The loss function matters."},
                    {"id": 2, "start": 9.0, "end": 12.0, "text": "Gradient Descent again."},
                ],
            }
        )

    def transcribe(self, audio_path: Path, language: str = "en") -> Any:
        self.calls.append((audio_path, language))
        return self.result


@pytest.fixture
def lecture_pdf() -> bytes:
    return make_text_pdf(LECTURE_LINES)


@pytest.fixture
def blank_pdf() -> bytes:
    return make_blank_pdf()


@pytest.fixture
def encrypted_pdf() -> bytes:
    return make_blank_pdf(encrypt=True)


@pytest.fixture
def lecture_pptx() -> bytes:
    return make_pptx([LECTURE_LINES[:2], [LECTURE_LINES[2]]], table=[["Loss Function", "Error"]])


@pytest.fixture
def fake_transcriber() -> FakeTranscriber:
    return FakeTranscriber()


@pytest.fixture
def settings(tmp_path: Path) -> CoreSettings:
    return CoreSettings(data_dir=tmp_path, max_upload_mb=1)


@pytest.fixture
def client(settings: CoreSettings, fake_transcriber: FakeTranscriber) -> Iterator[TestClient]:
    app = FastAPI()
    app.include_router(router, prefix="/api/c1")
    matcher = SemanticMatcher(enabled=False)
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[get_transcriber] = lambda: fake_transcriber
    app.dependency_overrides[get_matcher] = lambda: matcher
    with TestClient(app) as test_client:
        yield test_client
